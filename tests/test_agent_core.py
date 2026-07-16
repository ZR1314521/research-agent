from __future__ import annotations

import io
import json
import shutil
import unittest
from pathlib import Path
from typing import Any
from unittest.mock import Mock, patch

from research_agent.config import AgentConfig
from research_agent.context import ContextManager
from research_agent.core.agent import AgentLoop
from research_agent.core.contracts import ContractError, public_artifacts, refresh_artifact_profiles, register_artifacts, validate_result_quality
from research_agent.core.prompt_runtime import PromptRuntime
from research_agent.chat import ResearchChatAgent
from research_agent.capabilities.files import FileService
from research_agent.capabilities.literature import LiteratureService
from research_agent.capabilities.references import ReferenceService
from research_agent.capabilities.rag import RagService
from research_agent.capabilities.writing import WritingService
from research_agent.session import ChatSession, SessionStore
from research_agent.skill_registry import SkillRegistry
from research_agent.tools.llm_client import LLMResult
from research_agent.tools.llm_client import LLMClient
from research_agent.logging import ModelCallLogger


ROOT = Path(__file__).resolve().parents[1]


def text_response(text: str) -> LLMResult:
    return LLMResult(
        text, "test", "test", True, outcome="text",
        assistant_message={"role": "assistant", "content": text},
    )


def tool_response(name: str, arguments: dict, call_id: str = "call-1") -> LLMResult:
    calls = [{
        "id": call_id,
        "type": "function",
        "function": {"name": name, "arguments": json.dumps(arguments, ensure_ascii=False)},
    }]
    return LLMResult(
        "", "test", "test", True, outcome="native_tool_call",
        tool_calls=calls,
        assistant_message={"role": "assistant", "content": None, "tool_calls": calls},
    )


class AgentCoreTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = ROOT / ".test_runs" / self.id().replace(".", "_")
        shutil.rmtree(self.tmp, ignore_errors=True)
        self.tmp.mkdir(parents=True, exist_ok=True)
        self.config = AgentConfig.load(ROOT)
        object.__setattr__(self.config, "runs_dir", self.tmp / "runs")
        object.__setattr__(self.config, "llm_api_key", "test-key")
        self.registry = SkillRegistry(self.config.skills_dir)
        self.sessions = SessionStore(self.config.runs_dir)
        self.session = self.sessions.create()

    def tearDown(self) -> None:
        self.sessions.close()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_model_decides_conversation_without_hardcoded_reply(self) -> None:
        loop = AgentLoop(
            self.config,
            self.registry,
            self.sessions,
            self.sessions.directory(self.session.session_id),
        )
        with patch.object(
            loop.client,
            "chat",
            return_value=text_response("我是科研智能体，可以帮你处理文献和实验数据。"),
        ):
            result = loop.run(self.session, "你是谁")
        self.assertIn("科研智能体", result.message)
        self.assertEqual(result.skill, "")
        self.assertEqual(self.session.events[-1]["event"], "finished")

    def test_chat_uses_agent_loop_when_model_is_configured(self) -> None:
        agent = ResearchChatAgent(self.config)
        self.addCleanup(agent.close)
        session = agent.new_session()
        with patch(
            "research_agent.tools.llm_client.LLMClient.chat",
            return_value=text_response("模型驱动的自然回答。"),
        ):
            response = agent.handle(session, "你是谁")
        self.assertEqual(response.message, "模型驱动的自然回答。")
        self.assertEqual(response.skill, "")

    def test_tool_observation_is_replanned(self) -> None:
        loop = AgentLoop(
            self.config,
            self.registry,
            self.sessions,
            self.sessions.directory(self.session.session_id),
        )
        calls = iter(
            [
                tool_response("workspace-files", {"operation": "list", "path": "."}),
                text_response("我检查完了，当前没有失败请求。"),
            ]
        )
        loop.executor.execute = lambda tool, arguments, session: {
            "message": "请求 0 次，暂无失败。",
            "artifacts": {},
            "data": {"request_count": 0},
        }
        with patch.object(loop.client, "chat", side_effect=lambda *args, **kwargs: next(calls)) as mocked:
            result = loop.run(self.session, "刚才请求了几次")
        self.assertIn("检查完了", result.message)
        self.assertEqual(mocked.call_count, 2)
        self.assertIn("tool_observed", [event["event"] for event in self.session.events])
        self.assertEqual(self.session.model_messages[-2]["role"], "tool")

    def test_literature_batch_retrieves_all_queries_but_screens_once(self) -> None:
        service = LiteratureService(self.config, self.sessions.directory(self.session.session_id))
        client = Mock()
        client.search.side_effect = lambda query, **_kwargs: [
            {"title": f"Paper for {query}", "year": 2025, "abstract": "motor imagery EEG method", "venue": "Test"}
        ]
        service._client = lambda _source: client
        service.screen_with_edge = Mock(side_effect=lambda papers, *_args: (papers, [], []))

        result = service.search({
            "queries": ["EEG MI CNN", "EEG MI Transformer", "EEG MI GNN"],
            "sources": ["openalex"],
            "limit": 10,
            "request": "比较近三年 EEG MI 方法趋势",
        })

        self.assertEqual(client.search.call_count, 3)
        self.assertEqual(service.screen_with_edge.call_count, 1)
        self.assertEqual(result["data"]["count"], 3)
        self.assertEqual(result["progress"]["metrics"]["deduplicated"], 3)
        self.assertTrue(result["model_data"]["answer_ready"])
        self.assertEqual(len(result["model_data"]["papers"]), 3)
        self.assertLessEqual(
            ContextManager.estimate_tokens(result["model_data"]["papers"]),
            int(self.config.context_observation_budget * 0.72),
        )

    def test_literature_batch_bounds_screening_and_interleaves_queries(self) -> None:
        service = LiteratureService(self.config, self.sessions.directory(self.session.session_id))
        client = Mock()
        client.search.side_effect = lambda query, **_kwargs: [
            {"title": f"{query} result {rank}", "year": 2025, "abstract": "evidence"}
            for rank in range(4)
        ]
        service._client = lambda _source: client
        service.screen_with_edge = Mock(side_effect=lambda papers, *_args: (papers, [], []))

        service.search({
            "queries": ["query one", "query two", "query three"],
            "sources": ["openalex"],
            "limit": 3,
            "request": "compare methods",
        })

        evaluated = service.screen_with_edge.call_args.args[0]
        self.assertEqual(len(evaluated), 3)
        self.assertEqual(
            [paper["title"] for paper in evaluated],
            ["query one result 0", "query two result 0", "query three result 0"],
        )

    def test_literature_batch_applies_configured_query_and_screening_boundaries(self) -> None:
        object.__setattr__(self.config, "literature_max_batch_queries", 2)
        object.__setattr__(self.config, "literature_screening_limit", 3)
        service = LiteratureService(self.config, self.sessions.directory(self.session.session_id))
        client = Mock()
        client.search.side_effect = lambda query, **_kwargs: [
            {"title": f"{query} result {rank}", "year": 2025, "abstract": "evidence"}
            for rank in range(5)
        ]
        service._client = lambda _source: client
        service.screen_with_edge = Mock(side_effect=lambda papers, *_args: (papers, [], []))

        result = service.search({
            "queries": ["one", "two", "three", "four"],
            "sources": ["openalex"],
            "limit": 50,
            "request": "compare methods",
        })

        self.assertEqual(client.search.call_count, 2)
        self.assertEqual(len(service.screen_with_edge.call_args.args[0]), 3)
        self.assertEqual(result["progress"]["metrics"]["evaluated"], 3)

    def test_declarative_single_batch_policy_skips_duplicate_tool_calls(self) -> None:
        loop = AgentLoop(
            self.config,
            self.registry,
            self.sessions,
            self.sessions.directory(self.session.session_id),
        )
        tool_calls = [
            {
                "id": f"call-{index}",
                "type": "function",
                "function": {
                    "name": "academic-search-multisource",
                    "arguments": json.dumps({"query": query}),
                },
            }
            for index, query in enumerate(("first query", "second query"), 1)
        ]
        responses = iter([
            LLMResult(
                "", "test", "test", True, outcome="native_tool_call",
                tool_calls=tool_calls,
                assistant_message={"role": "assistant", "content": None, "tool_calls": tool_calls},
            ),
            text_response("我会根据已获得的证据直接回答。"),
        ])
        executed: list[str] = []
        captured: list[list[dict[str, Any]]] = []
        offered_tools: list[Any] = []

        def chat(_operation, messages, **kwargs):
            captured.append(json.loads(json.dumps(messages, ensure_ascii=False)))
            offered_tools.append(kwargs.get("tools"))
            return next(responses)

        loop.executor.execute = lambda tool, _arguments, _session: (
            executed.append(tool)
            or {
                "message": "找到足够证据。", "artifacts": {},
                "data": {}, "model_data": {"answer_ready": True},
            }
        )
        with patch.object(loop.client, "chat", side_effect=chat):
            result = loop.run(self.session, "研究这个问题")

        self.assertEqual(result.message, "我会根据已获得的证据直接回答。")
        self.assertEqual(executed, ["academic-search-multisource"])
        tool_payloads = [json.loads(item["content"]) for item in captured[-1] if item.get("role") == "tool"]
        self.assertEqual(len(tool_payloads), 2)
        self.assertEqual(tool_payloads[-1]["error_code"], "batch_policy")
        self.assertTrue(offered_tools[0])
        self.assertIsNone(offered_tools[1])

    def test_large_tool_observation_is_compact_even_when_total_context_has_room(self) -> None:
        manager = ContextManager(
            self.sessions.directory(self.session.session_id),
            context_window=1_000_000,
            context_budget=200_000,
            observation_budget=4_000,
        )
        observation = manager.observation(
            tool="academic-search-multisource",
            message="找到 12 篇论文。",
            data={"papers": [{"abstract": "evidence " * 50_000}]},
            artifacts={"active_papers": "screened_papers.json"},
            messages=[{"role": "user", "content": "分析研究趋势"}],
        )

        self.assertNotIn("data", observation)
        self.assertIn("data_ref", observation)
        self.assertLessEqual(manager.estimate_tokens(observation), 4_000)
        self.assertTrue(Path(observation["data_ref"]).exists())

    def test_compact_tool_observation_preserves_generic_answer_readiness(self) -> None:
        manager = ContextManager(
            self.sessions.directory(self.session.session_id),
            context_budget=20_000,
            observation_budget=700,
        )
        observation = manager.observation(
            tool="generic-research-tool",
            message="证据已整理。",
            data={
                "answer_ready": True,
                "missing_evidence": [],
                "completion_guidance": "Answer the user from the supplied evidence.",
                "records": [{"text": "evidence " * 10_000}],
            },
            artifacts={"support": "evidence.json"},
            messages=[{"role": "user", "content": "分析趋势"}],
        )

        self.assertTrue(observation["answer_ready"])
        self.assertEqual(observation["missing_evidence"], [])
        self.assertIn("completion_guidance", observation)
        self.assertNotIn("data", observation)

    def test_request_budget_includes_fixed_prompt_and_output_reserve(self) -> None:
        manager = ContextManager(
            self.sessions.directory(self.session.session_id),
            context_window=10_000,
            context_budget=1_000,
        )
        messages = [
            {"role": "user", "content": "分析这些结果"},
            {
                "role": "tool",
                "tool_call_id": "call-large",
                "content": json.dumps({"message": "完成", "data": {"raw": "x" * 2_400}}),
            },
        ]
        fixed = {"system": "policy " * 180, "tools": [{"schema": "y" * 600}]}

        fitted = manager.fit_to_budget(
            messages,
            fixed_context=fixed,
            reserve_tokens=200,
        )

        payload = json.loads(fitted[-1]["content"])
        self.assertNotIn("data", payload)
        self.assertLessEqual(
            manager.estimate_tokens(fitted) + manager.estimate_tokens(fixed) + 200,
            1_000,
        )

    def test_fallback_token_estimate_is_conservative_for_json_payloads(self) -> None:
        english_payload = {"abstracts": "x" * 367_000}
        chinese_payload = {"摘要": "抑郁症脑电研究" * 10_000}

        self.assertGreaterEqual(ContextManager.estimate_tokens(english_payload), 120_000)
        self.assertGreaterEqual(ContextManager.estimate_tokens(chinese_payload), 70_000)

    def test_fitted_cjk_text_respects_the_configured_budget(self) -> None:
        manager = ContextManager(
            self.sessions.directory(self.session.session_id),
            context_budget=1_000,
        )

        fitted = manager.fit_text("抑郁症脑电研究" * 10_000, label="cjk-regression")

        self.assertLessEqual(manager.estimate_tokens(fitted), 1_000)
        self.assertIn("Complete source:", fitted)

    def test_large_failed_tool_payload_is_externalized_too(self) -> None:
        manager = ContextManager(
            self.sessions.directory(self.session.session_id),
            context_budget=32_000,
            observation_budget=1_000,
        )
        payload = {
            "ok": False,
            "tool": "web-search",
            "error": "remote trace " * 20_000,
            "error_code": "remote_failure",
        }

        compact = manager.tool_payload(
            payload,
            messages=[{"role": "user", "content": "search"}],
        )

        self.assertEqual(compact["ok"], False)
        self.assertEqual(compact["error_code"], "remote_failure")
        self.assertIn("data_ref", compact)
        self.assertLessEqual(manager.estimate_tokens(compact), 1_000)

    def test_exploratory_search_ranks_low_score_papers_without_hard_excluding_them(self) -> None:
        service = LiteratureService(
            self.config,
            self.sessions.directory(self.session.session_id),
        )
        papers = [{
            "title": "EEG representation learning in clinical cohorts",
            "abstract": "A small exploratory neural representation study.",
            "year": 2024,
            "venue": "Clinical Neuroinformatics",
            "doi": "10.1/exploratory",
        }]

        assessment = {"papers": [{"id": "p0", "decision": "include", "score": 22, "reasons": ["exploratory but relevant"]}]}
        with patch.object(service.client, "complete", return_value=text_response(json.dumps(assessment))):
            included, excluded, edge = service.screen_with_edge(
                papers,
                "EEG MDD Transformer",
                2020,
                2026,
                [],
                [],
                [],
                False,
            )

        self.assertEqual(len(included), 1)
        self.assertFalse(excluded)
        self.assertFalse(edge)

    def test_explicit_screening_criteria_are_not_vetoed_by_a_score_threshold(self) -> None:
        service = LiteratureService(
            self.config,
            self.sessions.directory(self.session.session_id),
        )
        paper = {
            "title": "Clinical cohort representation study",
            "abstract": "A small clinical investigation.",
            "year": 2024,
            "keywords": [
                {"display_name": "EEG"},
                {"display_name": "MDD"},
                {"display_name": "Transformer"},
            ],
        }

        assessment = {
            "papers": [{
                "id": "p0", "decision": "include", "score": 80,
                "reasons": ["all concepts supported by supplied keywords"],
                "matched_requirements": ["EEG", "MDD", "Transformer"],
                "missing_requirements": [],
            }]
        }
        with patch.object(service.client, "complete", return_value=text_response(json.dumps(assessment))):
            included, excluded, edge = service.screen_with_edge(
                [paper],
                "EEG MDD Transformer",
                2020,
                2026,
                [],
                ["EEG", "MDD", "Transformer"],
                [],
                True,
            )

        self.assertEqual(len(included), 1)
        self.assertFalse(excluded)
        self.assertFalse(edge)

    def test_semantic_screening_requests_compact_per_paper_results(self) -> None:
        service = LiteratureService(
            self.config,
            self.sessions.directory(self.session.session_id),
        )
        response = text_response(json.dumps({
            "included": [{"id": "p0", "score": 81, "reason": "Relevant method and population."}],
            "edge": [],
            "excluded_ids": [],
            "next_query": "",
            "stop": True,
            "stop_reason": "sufficient",
        }))
        with patch.object(service.client, "complete", return_value=response) as complete:
            included, _excluded, _edge = service.screen_with_edge(
                [{"title": "A study", "abstract": "Relevant evidence", "year": 2024}],
                "a research question",
                None,
                None,
                [],
                [],
                [],
                False,
            )

        contract = json.loads(complete.call_args.args[1])["output_contract"]
        self.assertEqual(set(contract) & {"included", "edge", "excluded_ids"}, {"included", "edge", "excluded_ids"})
        self.assertIn("reason", contract["included"][0])
        self.assertNotIn("papers", contract)
        self.assertEqual(included[0]["screening_reasons"], ["Relevant method and population."])
        self.assertEqual(service.client.config.llm_timeout_seconds, self.config.structured_llm_timeout_seconds)
        self.assertEqual(service.client.config.llm_max_tokens, self.config.structured_llm_max_tokens)
        self.assertEqual(service.client.config.llm_retry, self.config.structured_llm_retry)

    def test_literature_semantics_follow_structured_evaluation_not_lexical_rules(self) -> None:
        service = LiteratureService(self.config, self.sessions.directory(self.session.session_id))
        papers = [
            {"title": "Depression EEG cohort", "abstract": "EEG diagnosis in patients with depression.", "year": 2024},
            {"title": "Depression EEG cohort", "abstract": "EEG study without a depression cohort.", "year": 2024},
        ]
        assessment = {
            "papers": [
                {"id": "p0", "decision": "include", "score": 94, "reasons": ["target population present"]},
                {"id": "p1", "decision": "exclude", "score": 5, "reasons": ["negated target population"], "exclusion_reason": "population_not_present"},
            ],
            "next_query": "depression EEG diagnostic biomarkers clinical cohort",
            "stop": False,
        }
        with patch.object(service.client, "complete", return_value=text_response(json.dumps(assessment))) as complete:
            included, excluded, edge = service.screen_with_edge(
                papers, "depression EEG", 2020, 2026, [], ["depression", "EEG"], [], True
            )

        self.assertEqual([paper["abstract"] for paper in included], [papers[0]["abstract"]])
        self.assertEqual(excluded[0]["exclusion_reason"], "population_not_present")
        self.assertFalse(edge)
        request = json.loads(complete.call_args.args[1])
        self.assertEqual(request["required_concepts"], ["depression", "EEG"])
        self.assertEqual(service._last_next_query, "depression EEG diagnostic biomarkers clinical cohort")

    def test_matrix_extraction_does_not_guess_semantic_fields_when_model_output_is_invalid(self) -> None:
        service = WritingService(self.config, self.sessions.directory(self.session.session_id))
        papers = self.tmp / "strict-evidence-papers.json"
        papers.write_text(json.dumps([{
            "title": "A proposed transformer for EEG",
            "abstract": "We propose a transformer and achieve 99% accuracy on a private cohort.",
            "year": 2024,
        }]), encoding="utf-8")

        with patch.object(service.client, "complete", return_value=text_response("not valid json")):
            result = service.summarize_papers(str(papers), {})

        row = result["data"]["rows"][0]
        self.assertEqual(row["method"], "NOT_REPORTED")
        self.assertEqual(row["innovation"], "NOT_REPORTED")
        self.assertEqual(row["key_findings"], "NOT_REPORTED")
        self.assertEqual(row["extraction_status"], "unavailable")

    def test_matrix_extraction_accepts_only_values_with_verifiable_evidence_quotes(self) -> None:
        service = WritingService(self.config, self.sessions.directory(self.session.session_id))
        abstract = "We evaluated a transformer on 120 participants. Accuracy improved by 4%."
        papers = self.tmp / "quoted-evidence-papers.json"
        papers.write_text(json.dumps([{"title": "Quoted evidence", "abstract": abstract, "year": 2024}]), encoding="utf-8")
        extraction = [{
            "index": 1,
            "method": {"status": "reported", "value": "transformer", "evidence_quote": "evaluated a transformer"},
            "dataset": {"status": "reported", "value": "999 participants", "evidence_quote": "999 participants"},
            "key_findings": {"status": "reported", "value": "Accuracy improved by 4%", "evidence_quote": "Accuracy improved by 4%"},
        }]

        with patch.object(service.client, "complete", return_value=text_response(json.dumps(extraction))):
            result = service.summarize_papers(str(papers), {})

        row = result["data"]["rows"][0]
        self.assertEqual(row["method"], "transformer")
        self.assertEqual(row["dataset"], "NOT_REPORTED")
        self.assertEqual(row["key_findings"], "Accuracy improved by 4%")
        self.assertTrue(result["model_data"]["answer_ready"])
        self.assertEqual(result["model_data"]["rows"][0]["method"], "transformer")
        evidence = json.loads(row["evidence_map"])
        self.assertEqual(evidence["method"]["quote"], "evaluated a transformer")
        self.assertEqual(evidence["dataset"]["status"], "invalid_evidence")

    def test_reference_enrichment_records_field_level_public_metadata_provenance(self) -> None:
        source = self.tmp / "incomplete-references.json"
        source.write_text(json.dumps([{
            "id": "ref-one", "type": "article-journal", "title": "A verified paper",
            "authors": [], "year": "", "venue": "", "doi": "10.1234/example",
        }]), encoding="utf-8")
        service = ReferenceService(self.config, self.sessions.directory(self.session.session_id))
        resolved = {
            "title": "A verified paper", "authors": ["A Author"], "year": "2024",
            "venue": "Journal of Verified Results", "doi": "10.1234/example",
            "volume": "12", "issue": "2", "pages": "10-20",
            "metadata_source": "crossref", "metadata_url": "https://api.crossref.org/works/10.1234/example",
        }
        with patch.object(service.metadata_client, "resolve", return_value=resolved):
            result = service.format({
                "path": str(source), "styles": ["gbt7714-numeric"],
                "output_mode": "list", "enrich_metadata": True,
            }, {})

        manifest = json.loads(Path(result["artifacts"]["reference_metadata_provenance"]).read_text(encoding="utf-8"))
        self.assertEqual(manifest["records"][0]["fields"]["year"]["source"], "crossref")
        normalized = json.loads(Path(result["artifacts"]["reference_input_normalized"]).read_text(encoding="utf-8"))
        self.assertEqual(normalized[0]["venue"], "Journal of Verified Results")

    def test_rag_answer_emits_a_complete_retrieval_manifest(self) -> None:
        note = self.tmp / "rag-source.txt"
        note.write_text("Accuracy increased across four experimental epochs.", encoding="utf-8")
        service = RagService(self.config, self.sessions.directory(self.session.session_id))
        with patch("research_agent.capabilities.rag.LLMClient.complete", return_value=text_response("Grounded answer.")):
            result = service.query({"query": "accuracy epochs", "scope": "uploaded"}, {"uploaded_note": str(note)})

        manifest_path = Path(result["artifacts"]["retrieval_manifest"])
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        self.assertEqual(manifest["query"], "accuracy epochs")
        self.assertEqual(manifest["inputs"], [str(note.resolve())])
        self.assertTrue(manifest["retrieved_chunks"])
        self.assertEqual(manifest["retrieved_chunks"][0]["source"], str(note.resolve()))

    def test_large_tool_payload_is_not_duplicated_in_session_state(self) -> None:
        loop = AgentLoop(
            self.config,
            self.registry,
            self.sessions,
            self.sessions.directory(self.session.session_id),
        )
        responses = iter([
            tool_response("academic-search-multisource", {"query": "EEG MDD"}),
            text_response("趋势分析完成。"),
        ])
        loop.executor.execute = lambda *_args: {
            "message": "找到 12 篇论文。",
            "artifacts": {"active_papers": str(self.tmp / "screened_papers.json")},
            "data": {"papers": [{"abstract": "large evidence " * 50_000}]},
        }

        requests: list[list[dict]] = []

        def chat(_operation, messages, **_kwargs):
            requests.append(json.loads(json.dumps(messages, ensure_ascii=False)))
            return next(responses)

        with patch.object(loop.client, "chat", side_effect=chat):
            result = loop.run(self.session, "分析 EEG MDD 方法趋势")

        self.assertEqual(result.message, "趋势分析完成。")
        self.assertNotIn("data", self.session.observations[-1])
        self.assertIn("data_ref", self.session.observations[-1])
        self.assertNotIn("papers", self.session.metadata["last_result"])
        self.assertTrue(Path(self.session.observations[-1]["data_ref"]).exists())
        model_tool_payload = json.loads(
            next(message["content"] for message in requests[1] if message["role"] == "tool")
        )
        self.assertNotIn("data", model_tool_payload)
        self.assertIn("data_ref", model_tool_payload)

    def test_existing_artifacts_do_not_force_network_confirmation(self) -> None:
        loop = AgentLoop(
            self.config,
            self.registry,
            self.sessions,
            self.sessions.directory(self.session.session_id),
        )
        self.session.artifacts["active_papers"] = str(self.tmp / "papers.json")
        executed: list[str] = []
        loop.executor.execute = lambda tool, arguments, session: (
            executed.append(tool)
            or {"message": "补充检索完成。", "artifacts": {}, "data": {}}
        )
        with patch.object(
            loop.client,
            "chat",
            side_effect=[
                tool_response("academic-search-multisource", {"query": "EEG"}),
                text_response("补充检索完成。"),
            ],
        ):
            result = loop.run(self.session, "补充一些文献")
        self.assertFalse(result.waiting)
        self.assertEqual(executed, ["academic-search-multisource"])

    def test_prompt_runtime_contains_reuse_first_policy(self) -> None:
        runtime = PromptRuntime(self.registry)
        prompt = runtime.system(self.session)
        self.assertIn("已有成果", prompt)
        self.assertIn("当前消息优先于旧任务", prompt)
        self.assertIn("普通说明、时间线、列表和分析不要放进代码块", prompt)
        self.assertIn("不得把趋势外推到未覆盖范围", prompt)
        self.assertTrue(runtime.tools_for_llm())

    def test_dynamic_session_state_does_not_change_system_prompt(self) -> None:
        runtime = PromptRuntime(self.registry)
        before = runtime.system(self.session)
        self.session.artifacts["active_papers"] = "screened_papers.json"

        after = runtime.system(self.session)
        dynamic = runtime.runtime_context(self.session)

        self.assertEqual(before, after)
        self.assertIn("active_papers", dynamic)

    def test_plan_safety_is_declared_by_capabilities_not_handler_name_checks(self) -> None:
        for name in (
            "reference-format-gbt7714",
            "20-ml-paper-writing",
            "doc-coauthoring",
            "humanizer",
            "canvas-design",
            "run-code",
        ):
            self.assertTrue(self.registry.get(name).write_access, name)

    def test_unconfigured_model_does_not_echo_internal_prompt(self) -> None:
        object.__setattr__(self.config, "llm_api_key", "")
        client = LLMClient(self.config, ModelCallLogger(self.sessions.directory(self.session.session_id)))
        result = client.complete("test", "SECRET_INTERNAL_PROMPT", fallback="")
        self.assertFalse(result.used_remote_model)
        self.assertNotIn("SECRET_INTERNAL_PROMPT", result.text)

    def test_model_text_response_is_classified_as_text(self) -> None:
        payload = {"choices": [{"finish_reason": "stop", "message": {"content": "  usable answer  "}}]}
        client = LLMClient(self.config, ModelCallLogger(self.sessions.directory(self.session.session_id)))
        with patch("urllib.request.urlopen", return_value=io.BytesIO(json.dumps(payload).encode("utf-8"))):
            result = client.complete("test", "prompt", fallback="fallback")
        self.assertEqual(result.text, "usable answer")
        self.assertTrue(result.used_remote_model)
        self.assertEqual(result.outcome, "text")

    def test_model_reasoning_without_final_content_is_not_generic_empty_response(self) -> None:
        payload = {"choices": [{"finish_reason": "stop", "message": {"content": None, "reasoning_content": "private chain of thought"}}]}
        run_dir = self.sessions.directory(self.session.session_id)
        client = LLMClient(self.config, ModelCallLogger(run_dir))
        with patch("urllib.request.urlopen", return_value=io.BytesIO(json.dumps(payload).encode("utf-8"))):
            result = client.complete("test", "prompt", fallback="fallback")
        self.assertEqual(result.text, "fallback")
        self.assertEqual(result.error, "reasoning_incomplete")
        self.assertEqual(result.outcome, "reasoning_incomplete")
        log_text = (run_dir / "model_call_log.jsonl").read_text(encoding="utf-8")
        self.assertNotIn("private chain of thought", log_text)
        self.assertEqual(json.loads(log_text)["metadata"]["response_shape"]["reasoning_chars"], len("private chain of thought"))

    def test_model_length_exhaustion_is_not_generic_empty_response(self) -> None:
        payload = {"choices": [{"finish_reason": "length", "message": {"content": ""}}]}
        client = LLMClient(self.config, ModelCallLogger(self.sessions.directory(self.session.session_id)))
        with patch("urllib.request.urlopen", return_value=io.BytesIO(json.dumps(payload).encode("utf-8"))):
            result = client.complete("test", "prompt", fallback="fallback")
        self.assertEqual(result.text, "fallback")
        self.assertEqual(result.error, "output_exhausted")
        self.assertEqual(result.outcome, "output_exhausted")

    def test_model_native_tool_calls_are_not_parsed_as_text(self) -> None:
        payload = {"choices": [{"finish_reason": "tool_calls", "message": {"content": None, "tool_calls": [{"id": "call_1", "type": "function", "function": {"name": "search", "arguments": "{}"}}]}}]}
        client = LLMClient(self.config, ModelCallLogger(self.sessions.directory(self.session.session_id)))
        with patch("urllib.request.urlopen", return_value=io.BytesIO(json.dumps(payload).encode("utf-8"))):
            result = client.complete("test", "prompt", fallback="fallback")
        self.assertEqual(result.text, "")
        self.assertTrue(result.tool_calls)
        self.assertEqual(result.outcome, "native_tool_call")

    def test_model_tool_calls_take_precedence_over_mixed_content(self) -> None:
        payload = {"choices": [{"finish_reason": "tool_calls", "message": {"content": "Do not treat this as a final answer", "tool_calls": [{"id": "call_1", "type": "function", "function": {"name": "search", "arguments": "{}"}}]}}]}
        client = LLMClient(self.config, ModelCallLogger(self.sessions.directory(self.session.session_id)))
        with patch("urllib.request.urlopen", return_value=io.BytesIO(json.dumps(payload).encode("utf-8"))):
            result = client.complete("test", "prompt", fallback="fallback")
        self.assertTrue(result.tool_calls)
        self.assertEqual(result.outcome, "native_tool_call")

    def test_reasoning_incomplete_waits_for_a_retry_without_running_a_tool(self) -> None:
        loop = AgentLoop(self.config, self.registry, self.sessions, self.sessions.directory(self.session.session_id))
        loop.executor.execute = lambda *args: self.fail("tool must not run")
        with patch.object(
            loop.client,
            "chat",
            return_value=LLMResult("", "test", "test", False, "reasoning_incomplete", "agent_decision", 1, "reasoning_incomplete"),
        ):
            result = loop.run(self.session, "继续当前研究任务")
        self.assertTrue(result.waiting)
        self.assertIn("尚未生成最终回答", result.message)
        self.assertNotIn("API key", result.message)
        self.assertEqual(self.session.status, "waiting_user")
        self.assertEqual(self.session.pending_action, {
            "type": "retry_current_request",
            "outcome": "reasoning_incomplete",
        })
        self.assertEqual(self.session.events[-1]["outcome"], "reasoning_incomplete")

    def test_output_exhausted_waits_for_a_retry_without_running_a_tool(self) -> None:
        loop = AgentLoop(self.config, self.registry, self.sessions, self.sessions.directory(self.session.session_id))
        loop.executor.execute = lambda *args: self.fail("tool must not run")
        with patch.object(
            loop.client,
            "chat",
            return_value=LLMResult("", "test", "test", False, "output_exhausted", "agent_decision", 1, "output_exhausted"),
        ):
            result = loop.run(self.session, "继续当前研究任务")
        self.assertTrue(result.waiting)
        self.assertIn("尚未生成最终回答", result.message)
        self.assertNotIn("模型地址", result.message)
        self.assertEqual(self.session.status, "waiting_user")
        self.assertEqual(self.session.pending_action["type"], "retry_current_request")
        self.assertEqual(self.session.events[-1]["outcome"], "output_exhausted")

    def test_repair_reasoning_incomplete_waits_for_a_retry(self) -> None:
        loop = AgentLoop(self.config, self.registry, self.sessions, self.sessions.directory(self.session.session_id))
        with patch.object(
            loop.client, "chat",
            return_value=LLMResult("", "test", "test", False, "model_request_failed", "agent_turn", 1, "reasoning_incomplete"),
        ):
            result = loop.run(self.session, "继续当前研究任务")
        self.assertTrue(result.waiting)
        self.assertEqual(self.session.status, "waiting_user")
        self.assertEqual(self.session.pending_action["outcome"], "reasoning_incomplete")

    def test_repair_output_exhausted_waits_for_a_retry(self) -> None:
        loop = AgentLoop(self.config, self.registry, self.sessions, self.sessions.directory(self.session.session_id))
        with patch.object(
            loop.client, "chat",
            return_value=LLMResult("", "test", "test", False, "model_request_failed", "agent_turn", 1, "output_exhausted"),
        ):
            result = loop.run(self.session, "继续当前研究任务")
        self.assertTrue(result.waiting)
        self.assertEqual(self.session.status, "waiting_user")
        self.assertEqual(self.session.pending_action["outcome"], "output_exhausted")

    def test_model_outcome_is_authoritative_over_error(self) -> None:
        loop = AgentLoop(self.config, self.registry, self.sessions, self.sessions.directory(self.session.session_id))
        with patch.object(
            loop.client,
            "chat",
            return_value=LLMResult("", "test", "test", False, "TimeoutError: timed out", "agent_decision", 1, "reasoning_incomplete"),
        ):
            result = loop.run(self.session, "继续当前研究任务")
        self.assertTrue(result.waiting)
        self.assertIn("尚未生成最终回答", result.message)
        self.assertNotIn("API key", result.message)
        self.assertEqual(self.session.events[-1]["outcome"], "reasoning_incomplete")

    def test_native_tool_call_is_an_unsupported_provider_error_without_running_a_tool(self) -> None:
        loop = AgentLoop(self.config, self.registry, self.sessions, self.sessions.directory(self.session.session_id))
        loop.executor.execute = lambda *args: self.fail("tool must not run")
        with patch.object(
            loop.client,
            "chat",
            return_value=LLMResult("", "test", "test", False, "native_tool_call", "agent_decision", 1, "native_tool_call"),
        ):
            result = loop.run(self.session, "搜索论文")
        self.assertFalse(result.waiting)
        self.assertIn("模型服务当前不可用", result.message)
        self.assertEqual(self.session.pending_action, None)
        self.assertEqual(self.session.events[-1]["outcome"], "native_tool_call")

    def test_transport_failure_keeps_the_unavailable_model_message(self) -> None:
        loop = AgentLoop(self.config, self.registry, self.sessions, self.sessions.directory(self.session.session_id))
        loop.executor.execute = lambda *args: self.fail("tool must not run")
        with patch.object(
            loop.client,
            "chat",
            return_value=LLMResult("", "test", "test", False, "TimeoutError: timed out", "agent_decision", 1, ""),
        ):
            result = loop.run(self.session, "搜索论文")
        self.assertFalse(result.waiting)
        self.assertIn("模型服务当前不可用", result.message)
        self.assertEqual(self.session.events[-1]["error"], "TimeoutError: timed out")

    def test_file_upload_accepts_model_local_file_path_alias(self) -> None:
        source = self.tmp / "paper.docx"
        source.write_bytes(b"not-a-real-docx-but-a-local-upload")
        result = FileService(self.sessions.directory(self.session.session_id)).register(
            {"local_file_path": str(source)}
        )
        self.assertEqual(result["data"]["files"][0]["status"], "registered")

    def test_document_summary_prompt_carries_request_and_evidence_boundary(self) -> None:
        service = WritingService(self.config, self.sessions.directory(self.session.session_id))
        request = "请详细解释这篇文章的贡献"
        with patch.object(
            service.client,
            "complete",
            return_value=LLMResult("## 问题\n\n已说明。", "test", "test", True),
        ) as model:
            result = service.summarize_document("The paper evaluates a model.", {"request": request})

        prompt = model.call_args.args[1]
        self.assertIn(request, prompt)
        self.assertIn("禁止编造", prompt)
        self.assertIn("详细", prompt)
        self.assertTrue(Path(result["artifacts"]["document_summary"]).exists())

    def test_reference_service_extracts_references_from_docx(self) -> None:
        from docx import Document

        source = self.tmp / "references.docx"
        document = Document()
        document.add_heading("References", level=1)
        document.add_paragraph("Smith, J. (2024). EEG biomarkers for depression. Journal of EEG, 12, 1-10.")
        document.save(source)
        service = ReferenceService(self.config, self.sessions.directory(self.session.session_id))
        parsed = '[{"id":"smith2024","type":"article-journal","authors":["John Smith"],"title":"EEG biomarkers for depression","year":"2024","venue":"Journal of EEG","volume":"12","pages":"1-10"}]'
        before = source.read_bytes()
        with patch(
            "research_agent.tools.llm_client.LLMClient.complete",
            return_value=LLMResult(parsed, "test", "test", True),
        ) as model:
            result = service.format({"path": str(source), "styles": ["nature"]}, {})
        records = __import__("json").loads(
            Path(result["artifacts"]["reference_input_normalized"]).read_text(encoding="utf-8")
        )
        self.assertEqual(records[0]["title"], "EEG biomarkers for depression")
        self.assertTrue(Path(result["artifacts"]["references_nature"]).exists())
        self.assertEqual(source.read_bytes(), before)
        model.assert_not_called()

    def test_reference_colon_creates_full_nature_copy_and_skips_sections(self) -> None:
        from docx import Document

        source = self.tmp / "article.docx"
        document = Document()
        document.add_heading("1. Introduction", level=1)
        document.add_paragraph("This sentence cites [1].")
        document.add_heading("Reference:", level=1)
        document.add_paragraph("Li A, Wang B. EEG decoding with a compact model[J]. Journal of EEG, 2024, 12(1): 1-9.")
        document.add_paragraph("Chen C, Zhao D. Transformer EEG classification[J]. Neural Systems, 2023, 8(2): 10-20.")
        document.save(source)
        before = source.read_bytes()
        result = ReferenceService(self.config, self.sessions.directory(self.session.session_id)).format(
            {"path": str(source), "styles": ["nature"], "convert_in_text": True}, {}
        )
        output = Path(result["artifacts"]["nature_document"])
        self.assertTrue(output.exists())
        self.assertEqual(source.read_bytes(), before)
        self.assertEqual(result["data"]["reference_count"], 2)
        self.assertEqual(result["data"]["in_text_citations"]["converted"], 1)
        rendered = Path(result["artifacts"]["references_nature"]).read_text(encoding="utf-8")
        self.assertNotIn("NEEDS_CHECK", rendered)
        self.assertIn("Li, A.", rendered)
        text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
        self.assertIn("References", text)
        self.assertNotIn("1. Introduction", Path(result["artifacts"]["reference_input_normalized"]).read_text(encoding="utf-8"))

    def test_numbered_sections_are_not_a_reference_fallback(self) -> None:
        from docx import Document

        source = self.tmp / "no_references.docx"
        document = Document()
        for title in ("1. Introduction", "2. Methods", "3. Conclusion"):
            document.add_heading(title, level=1)
        document.save(source)
        with self.assertRaisesRegex(ValueError, "No reference section"):
            ReferenceService(self.config, self.sessions.directory(self.session.session_id)).format(
                {"path": str(source), "styles": ["nature"]}, {}
            )

    def test_quality_boundary_rejects_missing_artifact(self) -> None:
        with self.assertRaisesRegex(ContractError, "missing artifacts"):
            validate_result_quality({"artifacts": {"report": str(self.tmp / "missing.md")}, "data": {}})

    def test_quality_boundary_accepts_a_valid_empty_search_result(self) -> None:
        pool = self.tmp / "active_papers.json"
        pool.write_text("[]", encoding="utf-8")
        validate_result_quality({"artifacts": {"active_papers": str(pool)}, "data": {"count": 0}})

    def test_task_plan_is_model_generated_and_revisable(self) -> None:
        loop = AgentLoop(self.config, self.registry, self.sessions, self.sessions.directory(self.session.session_id))
        self.session.status = "planning"
        with patch.object(loop.client, "chat", return_value=text_response("1. 读取文档\n2. 转换并验收")):
            result = loop.run(self.session, "convert this document")
        self.assertTrue(result.waiting)
        self.assertEqual(self.session.pending_action["type"], "plan_approval")
        self.assertIn("读取文档", result.message)

    def test_contract_rejects_ignored_parameters_and_replans(self) -> None:
        loop = AgentLoop(self.config, self.registry, self.sessions, self.sessions.directory(self.session.session_id))
        calls = iter([
            tool_response("doc-coauthoring", {"source": "latest_document", "instruction": "extract references"}),
            text_response("该工具不能提取参考文献；我需要改用引用格式化工具。"),
        ])
        with patch.object(loop.client, "chat", side_effect=lambda *args, **kwargs: next(calls)):
            result = loop.run(self.session, "为什么刚才失败")
        self.assertIn("不能提取参考文献", result.message)
        observed = [event for event in self.session.events if event["event"] == "tool_observed"][-1]
        self.assertEqual(observed["error_code"], "unknown_argument")

    def test_artifact_ledger_uses_types_not_only_paths(self) -> None:
        spec = self.registry.get("file-upload-router")
        register_artifacts(self.session, spec, {"latest_document": str(self.tmp / "paper.docx")})
        self.assertEqual(self.session.artifact_records["latest_document"]["type"], "WordDocument")
        self.assertEqual(self.session.artifact_records["latest_document"]["producer"], "file-upload-router")

    def test_artifact_ledger_exposes_only_explicit_user_outputs(self) -> None:
        spec = self.registry.get("academic-search-multisource")
        register_artifacts(
            self.session,
            spec,
            {
                "raw_papers": str(self.tmp / "papers_raw.json"),
                "paper_pool_markdown": str(self.tmp / "paper_pool.md"),
            },
        )

        visible = public_artifacts(self.session.artifact_records)

        self.assertNotIn("raw_papers", visible)
        self.assertEqual(visible["paper_pool_markdown"]["presentation"], "supporting")
        self.assertEqual(visible["paper_pool_markdown"]["label"], "筛选后的文献池")

    def test_old_artifact_records_are_upgraded_from_their_producer_contract(self) -> None:
        self.session.artifacts["paper_pool_markdown"] = str(self.tmp / "paper_pool.md")
        self.session.artifact_records["paper_pool_markdown"] = {
            "type": "PaperPool",
            "path": self.session.artifacts["paper_pool_markdown"],
            "producer": "literature-search-openalex",
        }

        changed = refresh_artifact_profiles(self.session, self.registry)

        self.assertTrue(changed)
        visible = public_artifacts(self.session.artifact_records)
        self.assertEqual(visible["paper_pool_markdown"]["presentation"], "supporting")

    def test_prompt_excludes_large_stale_result_payloads(self) -> None:
        self.session.metadata["last_result"] = {"huge": "x" * 50000}
        prompt = PromptRuntime(self.registry).system(self.session)
        self.assertNotIn("x" * 1000, prompt)

    def test_review_outline_returns_to_model_without_hidden_workflow(self) -> None:
        loop = AgentLoop(
            self.config,
            self.registry,
            self.sessions,
            self.sessions.directory(self.session.session_id),
        )
        loop.executor.execute = lambda tool, arguments, session: {
            "message": "综述大纲已生成。",
            "artifacts": {"review_framework": str(self.tmp / "review_framework.md")},
            "data": {"stage": "outline", "requires_confirmation": True},
        }
        with patch.object(
            loop.client,
            "chat",
            side_effect=[
                tool_response("systematic-literature-review", {"stage": "outline"}),
                text_response("综述大纲已生成，是否继续撰写由你决定。"),
            ],
        ):
            result = loop.run(self.session, "用已有论文写综述")
        self.assertFalse(result.waiting)
        self.assertIn("大纲已生成", result.message)
        self.assertIsNone(self.session.pending_action)


if __name__ == "__main__":
    unittest.main()
