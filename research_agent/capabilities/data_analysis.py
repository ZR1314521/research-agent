from __future__ import annotations

import csv
import json
import math
import statistics
from collections import defaultdict
from pathlib import Path
from typing import Any

from research_agent.capabilities.charting import ChartService


class ExperimentAnalysisService:
    def __init__(self, session_dir: Path):
        self.session_dir = session_dir

    def analyze(self, arguments: dict[str, Any]) -> dict[str, Any]:
        raw_path = str(arguments.get("path") or "").strip()
        if not raw_path:
            raise ValueError('请提供数据文件路径，例如：分析数据 "D:\\data\\experiment.xlsx"')
        path = Path(raw_path).expanduser().resolve()
        if not path.exists() or not path.is_file():
            raise FileNotFoundError(f"Data file not found: {path}")
        rows = self._load(path)
        if not rows:
            raise ValueError("Data file contains no rows")
        columns = list(rows[0])
        roles = self._column_roles(arguments.get("column_roles"), columns)
        id_like = roles["identifiers"]
        missing = {column: sum(self._blank(row.get(column)) for row in rows) for column in columns}
        numeric = self._numeric_columns(rows, columns, set(id_like))
        summaries, outliers = self._summaries(numeric)
        trends = self._trends(rows, numeric, arguments.get("trend_specs") or [])
        groups = self._groups(rows, numeric, arguments.get("group_specs") or [])
        chart_specs = self._chart_specs(arguments.get("chart_specs") or [], columns)
        chart_artifacts, chart_results = self._render_charts(path, chart_specs)
        schema = self._schema(rows, columns)
        payload = {
            "file": str(path),
            "row_count": len(rows),
            "columns": columns,
            "schema": schema,
            "column_roles": roles,
            "id_like_columns": id_like,
            "missing_values": missing,
            "numeric_summary": summaries,
            "outliers": outliers,
            "trends": trends,
            "group_comparisons": groups,
            "visualization_recommendations": chart_specs,
            "rendered_charts": chart_results,
            "assumptions": [
                "IQR outliers are exploratory flags and are not removed automatically.",
                "Column roles come only from the explicit analysis plan; names are not interpreted by regex.",
                "No trend or group comparison is inferred when its explicit specification is absent.",
            ],
        }
        summary_path = self.session_dir / "analysis_summary.json"
        plan_path = self.session_dir / "analysis_plan.json"
        report_path = self.session_dir / "analysis_report.md"
        outlier_path = self.session_dir / "outlier_flags.csv"
        summary_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        plan_path.write_text(json.dumps({
            "column_roles": roles,
            "trend_specs": arguments.get("trend_specs") or [],
            "group_specs": arguments.get("group_specs") or [],
            "chart_specs": chart_specs,
        }, ensure_ascii=False, indent=2), encoding="utf-8")
        report_path.write_text(self._report(payload), encoding="utf-8")
        with outlier_path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=["row", "column", "value", "lower_bound", "upper_bound"])
            writer.writeheader()
            writer.writerows(outliers)
        message = [f"已分析 {path.name}：{len(rows)} 行，{len(columns)} 列，{len(summaries)} 个数值列。"]
        message.append(f"缺失值共 {sum(missing.values())} 个，IQR 异常标记 {len(outliers)} 个。")
        for column, trend in list(trends.items())[:4]:
            message.append(f"- {column}: {trend['direction']}，斜率 {trend['slope']:.4g}")
        message.append(f"报告：{report_path}")
        return {
            "message": "\n".join(message),
            "artifacts": {
                "analysis_summary": str(summary_path),
                "analysis_report": str(report_path),
                "outlier_flags": str(outlier_path),
                "analysis_plan": str(plan_path),
                **chart_artifacts,
            },
            "data": payload,
            "model_data": {
                "answer_ready": True,
                "missing_evidence": [],
                "row_count": payload["row_count"],
                "columns": payload["columns"],
                "missing_values": payload["missing_values"],
                "numeric_summary": payload["numeric_summary"],
                "outliers": payload["outliers"],
                "trends": payload["trends"],
                "group_comparisons": payload["group_comparisons"],
                "completion_guidance": "Explain the verified statistics to the user; do not return raw JSON or code.",
            },
            "progress": {
                "summary": f"数据分析完成：{len(rows)} 行，{len(columns)} 列，标记 {len(outliers)} 个异常值",
                "metrics": {"rows": len(rows), "columns": len(columns), "outliers": len(outliers), "charts": len(chart_artifacts)},
            },
        }

    def _load(self, path: Path) -> list[dict[str, Any]]:
        suffix = path.suffix.lower()
        if suffix in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise RuntimeError("XLSX support requires openpyxl") from exc
            book = load_workbook(path, read_only=True, data_only=True)
            sheet = book.active
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                return []
            headers = [str(value or f"column_{index + 1}") for index, value in enumerate(values[0])]
            return [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]
        if suffix not in {".csv", ".tsv", ".txt"}:
            raise ValueError(f"Unsupported data format: {suffix}")
        sample = path.read_text(encoding="utf-8-sig", errors="ignore")[:4096]
        delimiter = "\t" if suffix == ".tsv" else ","
        if suffix == ".txt":
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
            except csv.Error:
                delimiter = "\t" if "\t" in sample else ","
        with path.open(newline="", encoding="utf-8-sig", errors="ignore") as handle:
            return list(csv.DictReader(handle, delimiter=delimiter))

    def _numeric_columns(
        self, rows: list[dict[str, Any]], columns: list[str], excluded: set[str]
    ) -> dict[str, list[tuple[int, float]]]:
        numeric: dict[str, list[tuple[int, float]]] = {}
        for column in columns:
            if column in excluded:
                continue
            values = []
            for index, row in enumerate(rows, 1):
                value = self._number(row.get(column))
                if value is not None:
                    values.append((index, value))
            nonblank = len(rows) - sum(self._blank(row.get(column)) for row in rows)
            if len(values) >= 2 and len(values) >= max(2, math.ceil(nonblank * 0.7)):
                numeric[column] = values
        return numeric

    def _summaries(self, numeric: dict[str, list[tuple[int, float]]]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        summaries: dict[str, Any] = {}
        flags: list[dict[str, Any]] = []
        for column, indexed in numeric.items():
            values = [value for _, value in indexed]
            q1, q3 = self._quantile(values, 0.25), self._quantile(values, 0.75)
            iqr = q3 - q1
            lower, upper = q1 - 1.5 * iqr, q3 + 1.5 * iqr
            for row, value in indexed:
                if value < lower or value > upper:
                    flags.append({"row": row, "column": column, "value": value, "lower_bound": lower, "upper_bound": upper})
            summaries[column] = {
                "count": len(values),
                "mean": statistics.mean(values),
                "median": statistics.median(values),
                "stdev": statistics.stdev(values) if len(values) > 1 else 0,
                "min": min(values),
                "q1": q1,
                "q3": q3,
                "max": max(values),
                "outlier_count_iqr": sum(1 for value in values if value < lower or value > upper),
            }
        return summaries, flags

    def _trends(
        self,
        rows: list[dict[str, Any]],
        numeric: dict[str, list[tuple[int, float]]],
        specs: list[dict[str, Any]],
    ) -> dict[str, Any]:
        trends: dict[str, Any] = {}
        for spec in specs:
            if not isinstance(spec, dict):
                continue
            x_column = str(spec.get("x") or "").strip()
            y_columns = [str(item) for item in spec.get("y") or []]
            if x_column not in rows[0]:
                raise ValueError(f"Trend x column does not exist: {x_column}")
            for column in y_columns:
                if column not in numeric:
                    raise ValueError(f"Trend y column is not numeric or was excluded: {column}")
                pairs = [
                    (self._number(row.get(x_column)), self._number(row.get(column))) for row in rows
                ]
                pairs = [(x, y) for x, y in pairs if x is not None and y is not None]
                if len(pairs) < 3:
                    continue
                x = [pair[0] for pair in pairs]
                y = [pair[1] for pair in pairs]
                x_mean, y_mean = statistics.mean(x), statistics.mean(y)
                denominator = sum((value - x_mean) ** 2 for value in x)
                slope = 0.0 if denominator == 0 else sum((a - x_mean) * (b - y_mean) for a, b in zip(x, y)) / denominator
                scale = max(y) - min(y)
                normalized = 0.0 if scale == 0 else slope * max(1, len(y) - 1) / scale
                direction = "increasing" if normalized > 0.1 else "decreasing" if normalized < -0.1 else "stable"
                trends[column] = {
                    "x": x_column, "slope": slope, "normalized_change": normalized, "direction": direction
                }
        return trends

    def _groups(
        self,
        rows: list[dict[str, Any]],
        numeric: dict[str, list[tuple[int, float]]],
        specs: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        results: list[dict[str, Any]] = []
        for spec in specs:
            if not isinstance(spec, dict):
                continue
            categorical = str(spec.get("group") or "").strip()
            if categorical not in rows[0]:
                raise ValueError(f"Group column does not exist: {categorical}")
            measures = [str(item) for item in spec.get("measures") or []]
            missing = [item for item in measures if item not in numeric]
            if missing:
                raise ValueError(f"Group measures are not numeric or were excluded: {', '.join(missing)}")
            result: dict[str, Any] = {"group_column": categorical, "means": {}}
            for numeric_column in measures:
                grouped: dict[str, list[float]] = defaultdict(list)
                for row in rows:
                    group = str(row.get(categorical) or "").strip()
                    value = self._number(row.get(numeric_column))
                    if group and value is not None:
                        grouped[group].append(value)
                result["means"][numeric_column] = {
                    group: {"count": len(values), "mean": statistics.mean(values)} for group, values in grouped.items()
                }
            results.append(result)
        return results

    def _schema(self, rows: list[dict[str, Any]], columns: list[str]) -> list[dict[str, Any]]:
        schema = []
        for column in columns:
            values = [row.get(column) for row in rows if not self._blank(row.get(column))]
            numeric_count = sum(self._number(value) is not None for value in values)
            schema.append({
                "name": column,
                "nonblank_count": len(values),
                "numeric_count": numeric_count,
                "numeric_ratio": numeric_count / len(values) if values else 0,
                "distinct_count": len({str(value) for value in values}),
                "sample": [str(value) for value in values[:5]],
            })
        return schema

    def _column_roles(self, raw: Any, columns: list[str]) -> dict[str, list[str]]:
        source = raw if isinstance(raw, dict) else {}
        roles = {
            role: list(dict.fromkeys(str(item) for item in source.get(role, []) if str(item).strip()))
            for role in ("identifiers", "order", "groups", "measures")
        }
        missing = [item for values in roles.values() for item in values if item not in columns]
        if missing:
            raise ValueError(f"Column roles reference missing columns: {', '.join(dict.fromkeys(missing))}")
        return roles

    def _chart_specs(self, raw: Any, columns: list[str]) -> list[dict[str, Any]]:
        if not isinstance(raw, list):
            raise ValueError("chart_specs must be a list")
        specs = []
        for item in raw:
            if not isinstance(item, dict):
                raise ValueError("Each chart specification must be an object")
            requested = [str(item.get("x") or ""), str(item.get("group") or ""), *[str(v) for v in item.get("y") or []]]
            missing = [value for value in requested if value and value not in columns]
            if missing:
                raise ValueError(f"Chart specification references missing columns: {', '.join(missing)}")
            specs.append(dict(item))
        return specs

    def _render_charts(self, path: Path, specs: list[dict[str, Any]]) -> tuple[dict[str, str], list[dict[str, Any]]]:
        service = ChartService(self.session_dir)
        artifacts: dict[str, str] = {}
        results: list[dict[str, Any]] = []
        for spec in specs:
            rendered = service.render({"path": str(path), **spec})
            artifacts.update(rendered.get("artifacts") or {})
            results.append(rendered.get("data") or {})
        return artifacts, results

    def _report(self, payload: dict[str, Any]) -> str:
        lines = [
            "# Experiment Data Analysis",
            "",
            f"File: `{payload['file']}`",
            f"Rows: {payload['row_count']}",
            "",
        ]
        if payload.get("id_like_columns"):
            lines.extend(["## Identifier Columns From Analysis Plan", ""])
            lines.extend(f"- `{column}`" for column in payload["id_like_columns"])
            lines.extend([""])
        lines.extend(["## Numeric Summary", ""])
        for column, item in payload["numeric_summary"].items():
            lines.append(
                f"- `{column}`: mean={item['mean']:.6g}, median={item['median']:.6g}, "
                f"sd={item['stdev']:.6g}, range=[{item['min']:.6g}, {item['max']:.6g}], "
                f"IQR outliers={item['outlier_count_iqr']}"
            )
        lines.extend(["", "## Missing Values", ""])
        lines.extend(f"- `{column}`: {count}" for column, count in payload["missing_values"].items())
        lines.extend(["", "## Trends", ""])
        lines.extend(
            f"- `{column}`: {item['direction']} (slope={item['slope']:.6g})" for column, item in payload["trends"].items()
        )
        lines.extend(["", "## Visualization Recommendations", ""])
        lines.extend(f"- {item}" for item in payload["visualization_recommendations"])
        return "\n".join(lines) + "\n"

    def _quantile(self, values: list[float], fraction: float) -> float:
        ordered = sorted(values)
        position = (len(ordered) - 1) * fraction
        low, high = math.floor(position), math.ceil(position)
        if low == high:
            return ordered[low]
        return ordered[low] * (high - position) + ordered[high] * (position - low)

    def _number(self, value: Any) -> float | None:
        try:
            return None if self._blank(value) else float(value)
        except (TypeError, ValueError):
            return None

    def _blank(self, value: Any) -> bool:
        return value is None or str(value).strip() == ""
