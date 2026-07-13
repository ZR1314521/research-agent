from __future__ import annotations

import csv
import json
import re
import time
from pathlib import Path
from typing import Any


class DataTransformService:
    def __init__(self, session_dir: Path):
        self.session_dir = session_dir

    def transform(self, arguments: dict[str, Any], artifacts: dict[str, str]) -> dict[str, Any]:
        raw_path = str(arguments.get("path") or artifacts.get("latest_data") or "").strip()
        if not raw_path:
            raise ValueError("请先上传数据文件，或在请求中提供 CSV/TSV/XLSX 路径")
        source = Path(raw_path).expanduser().resolve()
        if not source.exists():
            raise FileNotFoundError(f"Data file not found: {source}")

        rows = self._load(source)
        if not rows:
            raise ValueError("Data file contains no rows")
        columns = list(rows[0])
        ops = list(arguments.get("ops") or arguments.get("transform_ops") or [])
        keep_columns = [str(item) for item in arguments.get("keep_columns") or []]
        if not ops:
            ops = self._ops_from_request(str(arguments.get("request") or ""))
        if not ops:
            raise ValueError("没有识别到数据转换动作，例如归一化、筛掉缺失值、只保留这些列")

        transformed = [dict(row) for row in rows]
        log: list[dict[str, Any]] = []
        if "drop_missing" in ops:
            before = len(transformed)
            transformed = [row for row in transformed if all(str(value or "").strip() for value in row.values())]
            log.append({"operation": "drop_missing", "before": before, "after": len(transformed)})
        if "keep_columns" in ops:
            if not keep_columns:
                keep_columns = self._keep_columns_from_request(str(arguments.get("request") or ""), columns)
            missing = [column for column in keep_columns if column not in columns]
            if missing:
                raise ValueError(f"这些列不存在：{', '.join(missing)}")
            transformed = [{column: row.get(column, "") for column in keep_columns} for row in transformed]
            columns = keep_columns
            log.append({"operation": "keep_columns", "columns": keep_columns})
        for op in ("normalize", "standardize"):
            if op in ops:
                transformed, details = self._scale(transformed, op)
                log.append({"operation": op, "columns": details})

        output = self._output_path(arguments, source)
        output.parent.mkdir(parents=True, exist_ok=True)
        output = self._versioned(output)
        self._write_csv(output, transformed, list(transformed[0]) if transformed else columns)
        log_path = self.session_dir / "data_transform_log.json"
        log_path.write_text(
            json.dumps(
                {
                    "timestamp": time.strftime("%Y-%m-%dT%H:%M:%S"),
                    "source": str(source),
                    "output": str(output),
                    "operations": log,
                    "source_preserved": True,
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        return {
            "message": f"数据转换完成，原文件未覆盖。\n输出：{output}\n转换日志：{log_path}",
            "artifacts": {"transformed_data": str(output), "data_transform_log": str(log_path), "latest_data": str(output)},
            "data": {"output": str(output), "operations": log},
        }

    def _load(self, path: Path) -> list[dict[str, Any]]:
        suffix = path.suffix.lower()
        if suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            book = load_workbook(path, read_only=True, data_only=True)
            sheet = book.active
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                return []
            headers = [str(value or f"column_{index + 1}") for index, value in enumerate(values[0])]
            return [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]
        delimiter = "\t" if suffix == ".tsv" else ","
        if suffix == ".txt":
            sample = path.read_text(encoding="utf-8-sig", errors="ignore")[:4096]
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
            except csv.Error:
                delimiter = "\t" if "\t" in sample else ","
        if suffix not in {".csv", ".tsv", ".txt"}:
            raise ValueError(f"Unsupported data transform format: {suffix}")
        with path.open(newline="", encoding="utf-8-sig", errors="ignore") as handle:
            return list(csv.DictReader(handle, delimiter=delimiter))

    def _scale(self, rows: list[dict[str, Any]], op: str) -> tuple[list[dict[str, Any]], list[str]]:
        numeric = {}
        for column in rows[0]:
            values = []
            for row in rows:
                try:
                    values.append(float(row.get(column, "")))
                except (TypeError, ValueError):
                    values = []
                    break
            if len(values) == len(rows) and len(values) > 1 and not self._id_like(column):
                numeric[column] = values
        result = [dict(row) for row in rows]
        for column, values in numeric.items():
            low, high = min(values), max(values)
            mean = sum(values) / len(values)
            variance = sum((value - mean) ** 2 for value in values) / max(1, len(values) - 1)
            std = variance ** 0.5
            for row, value in zip(result, values):
                if op == "normalize":
                    row[column] = "0" if high == low else f"{(value - low) / (high - low):.10g}"
                else:
                    row[column] = "0" if std == 0 else f"{(value - mean) / std:.10g}"
        return result, list(numeric)

    def _id_like(self, column: str) -> bool:
        return bool(re.search(r"(^id$|subject|被试|编号|trial|epoch)", column, re.IGNORECASE))

    def _write_csv(self, path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=columns)
            writer.writeheader()
            writer.writerows(rows)

    def _output_path(self, arguments: dict[str, Any], source: Path) -> Path:
        raw = str(arguments.get("output_path") or "").strip()
        if raw:
            path = Path(raw).expanduser()
        else:
            path = self.session_dir / f"{source.stem}_transformed.csv"
        if path.suffix.lower() not in {".csv", ".txt"}:
            path = path.with_suffix(".csv")
        return path

    def _versioned(self, path: Path) -> Path:
        if not path.exists():
            return path
        for index in range(2, 1000):
            candidate = path.with_name(f"{path.stem}_v{index}{path.suffix}")
            if not candidate.exists():
                return candidate
        raise RuntimeError(f"Cannot create versioned path for {path}")

    def _ops_from_request(self, request: str) -> list[str]:
        ops = []
        if re.search(r"归一化|normalize", request, re.IGNORECASE):
            ops.append("normalize")
        if re.search(r"标准化|standardize", request, re.IGNORECASE):
            ops.append("standardize")
        if re.search(r"筛掉缺失|删除缺失|drop", request, re.IGNORECASE):
            ops.append("drop_missing")
        if re.search(r"只保留|保留这些列|keep", request, re.IGNORECASE):
            ops.append("keep_columns")
        return ops

    def _keep_columns_from_request(self, request: str, columns: list[str]) -> list[str]:
        quoted = re.findall(r"`([^`]+)`|\"([^\"]+)\"|'([^']+)'", request)
        values = [next(item for item in match if item) for match in quoted if any(match)]
        if values:
            return [value for value in values if value in columns]
        return [column for column in columns if column in request]

