from __future__ import annotations

import re
import hashlib
import json
import shutil
import subprocess
from zipfile import ZipFile
from pathlib import Path
from typing import Any


class DocumentService:
    def __init__(self, session_dir: Path):
        self.session_dir = session_dir

    def export_docx(self, arguments: dict[str, Any], artifacts: dict[str, str]) -> dict[str, Any]:
        if self._wants_empty(arguments):
            return self._empty_docx(arguments)
        source = self._source(arguments, artifacts)
        if not source:
            raise ValueError("没有找到可导出的 Markdown 或文本成果")
        try:
            from docx import Document
            from docx.shared import Pt
        except ImportError as exc:
            raise RuntimeError("DOCX output requires python-docx") from exc
        text = source.read_text(encoding="utf-8", errors="ignore")
        output = self._output_path(arguments, source)
        if output.suffix.lower() != ".docx":
            output = output.with_suffix(".docx")
        output.parent.mkdir(parents=True, exist_ok=True)
        output = self._versioned(output)
        document = Document()
        normal = document.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(10.5)
        for raw in text.splitlines():
            line = raw.rstrip()
            if not line:
                document.add_paragraph()
            elif line.startswith("### "):
                document.add_heading(line[4:], level=3)
            elif line.startswith("## "):
                document.add_heading(line[3:], level=2)
            elif line.startswith("# "):
                document.add_heading(line[2:], level=1)
            elif re.match(r"^[-*]\s+", line):
                document.add_paragraph(re.sub(r"^[-*]\s+", "", line), style="List Bullet")
            elif re.match(r"^\d+\.\s+", line):
                document.add_paragraph(re.sub(r"^\d+\.\s+", "", line), style="List Number")
            else:
                document.add_paragraph(self._plain(line))
        document.save(output)
        return {
            "message": f"Word 文档已生成：{output}",
            "artifacts": {"docx_output": str(output), "docx_source": str(source)},
            "data": {},
        }

    def convert(self, arguments: dict[str, Any], artifacts: dict[str, str]) -> dict[str, Any]:
        raw = str(arguments.get("path") or "").strip()
        source = Path(raw).expanduser().resolve() if raw else self._source(arguments, artifacts)
        if not source or not source.is_file() or source.suffix.lower() not in {".docx", ".md", ".markdown"}:
            raise ValueError("document-convert requires an existing DOCX or Markdown file")
        imports = self.session_dir / "artifacts" / "imports"
        imports.mkdir(parents=True, exist_ok=True)
        imported = imports / source.name
        if not imported.exists():
            shutil.copy2(source, imported)
        target_type = str(arguments.get("target_format") or ("markdown" if source.suffix.lower() == ".docx" else "docx")).lower()
        if target_type == "md":
            target_type = "markdown"
        requested = str(arguments.get("output_path") or "").strip()
        output = Path(requested).expanduser() if requested else self.session_dir / "artifacts" / f"{source.stem}.{ 'md' if target_type == 'markdown' else 'docx'}"
        if not output.is_absolute(): output = self.session_dir / "artifacts" / output
        output.parent.mkdir(parents=True, exist_ok=True)
        if output.exists() and not arguments.get("confirmed"):
            output = self._versioned(output)
        pandoc = self._pandoc()
        media = output.with_name(f"{output.stem}_media")
        command = [str(pandoc), str(imported), "-o", str(output)]
        if target_type == "markdown": command.extend(["-t", "gfm", f"--extract-media={media}"])
        elif target_type != "docx": raise ValueError("target_format must be markdown or docx")
        result = subprocess.run(command, capture_output=True, text=True, timeout=90)
        if result.returncode or not output.exists(): raise RuntimeError((result.stderr or result.stdout or "pandoc conversion failed").strip())
        report = output.with_name(f"{output.stem}_conversion_report.json")
        report.write_text(json.dumps({"source": str(source), "import_copy": str(imported), "output": str(output), "media_dir": str(media) if media.exists() else "", "source_sha256": self._sha256(source), "converter": "pandoc", "limitations": ["Markdown cannot preserve Word pagination, tracked changes, comments, or floating layout losslessly."]}, ensure_ascii=False, indent=2), encoding="utf-8")
        return {"message": f"Document converted: {output}\nConversion report: {report}", "artifacts": {"converted_document": str(output), "conversion_report": str(report), "imported_document": str(imported)}, "data": {"target_format": target_type}}

    def readable_text(self, path: str) -> str:
        source = Path(path).expanduser().resolve()
        if not source.is_file():
            raise ValueError("document-summary requires an existing file")
        suffix = source.suffix.lower()
        if suffix in {".md", ".markdown"}:
            return source.read_text(encoding="utf-8", errors="ignore").strip()
        if suffix == ".pdf":
            from markitdown import MarkItDown
            return MarkItDown().convert(str(source)).text_content.strip()
        if suffix == ".docx":
            try:
                from docx import Document
                return "\n".join(paragraph.text for paragraph in Document(source).paragraphs).strip()
            except Exception:
                with ZipFile(source) as archive:
                    xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
                return "\n".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", xml)).strip()
        raise ValueError(f"document-summary does not support {suffix} files")

    def _pandoc(self) -> Path:
        import os
        winget = Path(os.environ.get("LOCALAPPDATA", "")) / "Microsoft" / "WinGet" / "Packages"
        candidates = [shutil.which("pandoc"), str(Path(os.environ.get("ProgramFiles", "C:/Program Files")) / "Pandoc" / "pandoc.exe")]
        if winget.exists(): candidates.extend(str(path) for path in winget.glob("JohnMacFarlane.Pandoc_*/*/pandoc.exe"))
        for raw in candidates:
            if raw and Path(raw).exists(): return Path(raw)
        raise RuntimeError("Pandoc is required for document conversion")

    def _word_to_html(self, source: Path, output: Path) -> None:
        src, dst = str(source).replace("'", "''"), str(output).replace("'", "''")
        script = f"$ErrorActionPreference='Stop';$w=New-Object -ComObject Word.Application;$w.Visible=$false;try{{$d=$w.Documents.Open('{src}');$d.SaveAs([ref]'{dst}',10);$d.Close()}}finally{{$w.Quit()}}"
        result = subprocess.run(["powershell", "-NoProfile", "-Command", script], capture_output=True, text=True, timeout=45)
        if result.returncode or not output.exists():
            raise RuntimeError((result.stderr or result.stdout or "Microsoft Word HTML export failed").strip())

    def _sha256(self, path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""): digest.update(chunk)
        return digest.hexdigest()

    def _source(self, arguments: dict[str, Any], artifacts: dict[str, str]) -> Path | None:
        requested = str(arguments.get("path") or "").strip()
        if requested and Path(requested).exists():
            return Path(requested)
        for key in (
            "review_framework",
            "write_paper_section",
            "revise_document",
            "humanize_text",
            "analysis_report",
            "literature_matrix_md",
            "paper_pool_markdown",
        ):
            path = artifacts.get(key)
            if path and Path(path).exists() and Path(path).suffix.lower() in {".md", ".txt"}:
                return Path(path)
        return None

    def _plain(self, line: str) -> str:
        line = re.sub(r"\[([^]]+)]\(([^)]+)\)", r"\1 (\2)", line)
        return re.sub(r"[*_`]", "", line)

    def _wants_empty(self, arguments: dict[str, Any]) -> bool:
        request = str(arguments.get("request") or "")
        return bool(re.search(r"(新建|创建|建个).+\.docx|什么都别放|空\s*(?:word|docx|文档)", request, re.IGNORECASE))

    def _empty_docx(self, arguments: dict[str, Any]) -> dict[str, Any]:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX output requires python-docx") from exc
        output = self._output_path(arguments, None)
        output.parent.mkdir(parents=True, exist_ok=True)
        output = self._versioned(output)
        Document().save(output)
        return {
            "message": f"空 Word 文档已创建：{output}",
            "artifacts": {"docx_output": str(output)},
            "data": {"empty": True},
        }

    def _output_path(self, arguments: dict[str, Any], source: Path | None) -> Path:
        raw = str(arguments.get("output_path") or "").strip()
        if not raw:
            request = str(arguments.get("request") or "")
            quoted = re.search(r"[\"“”']([^\"“”']+\.docx)[\"“”']", request, re.IGNORECASE)
            simple = re.findall(r"([^\s\"“”']+\.docx)", request, re.IGNORECASE)
            raw = (quoted.group(1) if quoted else simple[-1] if simple else "").strip()
        if raw:
            path = Path(raw).expanduser()
            if not path.is_absolute():
                path = self.session_dir / path.name
            return path
        stem = source.stem if source else "test"
        return self.session_dir / f"{stem}.docx"

    def _versioned(self, path: Path) -> Path:
        if not path.exists():
            return path
        for index in range(2, 1000):
            candidate = path.with_name(f"{path.stem}_v{index}{path.suffix}")
            if not candidate.exists():
                return candidate
        raise RuntimeError(f"Cannot create versioned path for {path}")
