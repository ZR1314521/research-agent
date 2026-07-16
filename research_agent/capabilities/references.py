from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import sys
from copy import deepcopy
from pathlib import Path
from typing import Any

from research_agent.config import AgentConfig
from research_agent.tools.crossref import CrossrefClient


class ReferenceService:
    """Conservative reference parsing and formatting; never claims empty output is valid."""

    def __init__(self, config: AgentConfig, session_dir: Path):
        self.config = config
        self.session_dir = session_dir
        self.metadata_client = CrossrefClient(config)

    def format(self, arguments: dict[str, Any], artifacts: dict[str, str]) -> dict[str, Any]:
        raw_path = str(arguments.get("path") or artifacts.get("latest_references") or artifacts.get("active_papers") or artifacts.get("latest_document") or "").strip()
        if not raw_path:
            raise ValueError("Reference input is required")
        source = Path(raw_path).expanduser().resolve()
        if not source.exists():
            raise FileNotFoundError(f"Reference input not found: {source}")

        normalized, refs, reference_index = self._normalize_input(source)
        refs, provenance = self._enrich_metadata(refs, bool(arguments.get("enrich_metadata")))
        normalized.write_text(json.dumps(refs, ensure_ascii=False, indent=2), encoding="utf-8")
        provenance_path = self.session_dir / "reference_metadata_provenance.json"
        provenance_path.write_text(json.dumps(provenance, ensure_ascii=False, indent=2), encoding="utf-8")
        quality_report = self._quality_gate(source, refs, reference_index)
        profiles_path = self.config.rules_dir / "style_profiles.json"
        profiles = json.loads(profiles_path.read_text(encoding="utf-8"))
        supported = {item["key"] for item in profiles}
        styles = [self._style(str(item)) for item in arguments.get("styles") or ["gbt7714-numeric"]]
        styles = [style for style in dict.fromkeys(styles) if style in supported]
        if not styles:
            raise ValueError("No supported reference style was requested")
        command = [
            sys.executable,
            str(self.config.scripts_dir / "format_references_strict.py"),
            str(normalized),
            "--profiles", str(profiles_path),
            "--out-dir", str(self.session_dir),
        ]
        for style in styles:
            command.extend(["--style", style])
        result = subprocess.run(command, cwd=self.config.root_dir, text=True, capture_output=True)
        if result.returncode != 0:
            raise RuntimeError((result.stderr or result.stdout or "Reference formatter failed").strip())

        outputs = {
            "reference_input_normalized": str(normalized),
            "reference_quality_report": str(quality_report),
            "citation_check_report": str(self.session_dir / "citation_check_report.md"),
            "references_bib": str(self.session_dir / "references.bib"),
            "reference_metadata_provenance": str(provenance_path),
        }
        for style in styles:
            outputs[f"references_{style}"] = str(self.session_dir / f"references_{style.replace('-', '_')}.md")
        document_style = styles[-1] if styles else ""
        style_markdown = outputs.get(f"references_{document_style}")
        if style_markdown and "NEEDS_CHECK" in Path(style_markdown).read_text(encoding="utf-8", errors="ignore"):
            raise RuntimeError(f"Reference quality gate failed after formatting. Report: {quality_report}")

        data: dict[str, Any] = {
            "styles": styles,
            "reference_count": len(refs),
            "unresolved_count": 0,
            "metadata_enriched_records": sum(bool(item.get("enriched")) for item in provenance["records"]),
        }
        if source.suffix.lower() == ".docx" and style_markdown and str(arguments.get("output_mode") or "document") != "list":
            target, audit_path, source_hash = self._write_nature_document(
                source, refs, reference_index, Path(style_markdown), arguments, document_style
            )
            audit = json.loads(audit_path.read_text(encoding="utf-8"))
            outputs.update({"formatted_document": str(target), "citation_audit": str(audit_path)})
            if document_style == "nature":
                outputs["nature_document"] = str(target)
            data.update({"document_output": str(target), "source_sha256": source_hash, "in_text_citations": audit, "unresolved_count": len(audit["unresolved"])})

        lines = [f"已将 {len(refs)} 条参考文献转换为 {document_style} 格式。"]
        if outputs.get("formatted_document"):
            lines.append(f"格式化文档: {outputs['formatted_document']}")
        if outputs.get("citation_audit") and data.get("unresolved_count"):
            lines.append(f"注意: {data['unresolved_count']} 条文中引用未能匹配，详见 {outputs['citation_audit']}")
        return {
            "message": "\n".join(lines),
            "artifacts": outputs,
            "data": data,
            "progress": {
                "summary": f"参考文献处理完成：{len(refs)} 条，输出 {len(styles)} 种格式",
                "metrics": {
                    "references": len(refs),
                    "styles": len(styles),
                    "enriched": data["metadata_enriched_records"],
                    "unresolved": data["unresolved_count"],
                },
            },
        }

    def _enrich_metadata(
        self, refs: list[dict[str, Any]], enabled: bool
    ) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        bibliographic_fields = ("title", "authors", "year", "venue", "volume", "issue", "pages", "publisher", "doi", "url")
        enriched_refs: list[dict[str, Any]] = []
        records: list[dict[str, Any]] = []
        for index, raw in enumerate(refs, 1):
            reference = dict(raw)
            fields = {
                field: {"source": "input" if reference.get(field) else "missing", "value": reference.get(field) or ""}
                for field in bibliographic_fields
            }
            remote: dict[str, Any] | None = None
            error = ""
            if enabled and reference.get("doi"):
                try:
                    remote = self.metadata_client.resolve(reference)
                except Exception as exc:
                    error = f"{type(exc).__name__}: {exc}"
            changed = False
            if remote:
                source_name = str(remote.get("metadata_source") or "public_metadata")
                for field in bibliographic_fields:
                    if not reference.get(field) and remote.get(field):
                        reference[field] = remote[field]
                        fields[field] = {"source": source_name, "value": remote[field]}
                        changed = True
            enriched_refs.append(reference)
            records.append({
                "index": index,
                "id": reference.get("id") or "",
                "doi": reference.get("doi") or "",
                "enriched": changed,
                "match_basis": (remote or {}).get("match_basis") or "",
                "metadata_url": (remote or {}).get("metadata_url") or "",
                "error": error,
                "fields": fields,
            })
        return enriched_refs, {"requested": enabled, "records": records}

    def _normalize_input(self, source: Path) -> tuple[Path, list[dict[str, Any]], int | None]:
        suffix = source.suffix.lower()
        reference_index: int | None = None
        if suffix == ".json":
            payload = json.loads(source.read_text(encoding="utf-8-sig"))
            refs = payload if isinstance(payload, list) else payload.get("references") or payload.get("papers") or payload.get("results") or []
        elif suffix == ".bib":
            refs = self._bibtex(source.read_text(encoding="utf-8-sig", errors="ignore"))
        elif suffix in {".ris", ".nbib", ".enw"}:
            refs = self._tagged(source.read_text(encoding="utf-8-sig", errors="ignore"))
        elif suffix == ".docx":
            refs, reference_index = self._docx_references(source)
        else:
            raise ValueError(f"Unsupported reference input format: {suffix}")
        if not isinstance(refs, list):
            raise ValueError("Reference input must contain a list of records")
        path = self.session_dir / "references_input_normalized.json"
        path.write_text(json.dumps(refs, ensure_ascii=False, indent=2), encoding="utf-8")
        return path, [item for item in refs if isinstance(item, dict)], reference_index

    def _docx_references(self, source: Path) -> tuple[list[dict[str, Any]], int]:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX reference extraction requires python-docx") from exc
        paragraphs = [paragraph.text.strip() for paragraph in Document(source).paragraphs]
        reference_index, lines = self._reference_section(paragraphs)
        if not lines:
            raise ValueError("No reference section was recognized in the Word document")
        refs = [self._parse_reference_line(line, index) for index, line in enumerate(lines, 1)]
        refs = [item for item in refs if item.get("title") or item.get("year")]
        if not refs:
            raise ValueError("No bibliographic entries could be parsed from the reference section")
        return refs, reference_index

    def _reference_section(self, paragraphs: list[str]) -> tuple[int, list[str]]:
        heading = re.compile(r"^\s*(references?|bibliography|\u53c2\u8003\u6587\u732e)\s*[:\uff1a]?\s*$", re.IGNORECASE)
        for index, value in enumerate(paragraphs):
            if heading.fullmatch(value):
                return index, [item for item in paragraphs[index + 1 : index + 201] if item]
        numbered = [
            value for value in paragraphs
            if re.match(r"^\[?\d{1,3}[\].\u3002]\s+", value)
            and (re.search(r"\b(?:19|20)\d{2}\b", value) or re.search(r"\[[JCMDRSP]\]", value, re.IGNORECASE))
        ]
        return (-1, numbered[:200]) if numbered else (-1, [])

    def _parse_reference_line(self, raw: str, index: int) -> dict[str, Any]:
        text = re.sub(r"^\[?\d{1,3}[\].\u3002]\s*", "", raw).strip()
        marker = re.search(r"\[([JCMDRSP]|EB/OL)\]", text, re.IGNORECASE)
        if marker:
            before, after = text[: marker.start()].strip(), text[marker.end() :].lstrip(". /")
            parts = [item.strip() for item in before.rsplit(".", 1)]
            authors = self._source_order_authors(parts[0].split(",")) if parts else []
            title = parts[1] if len(parts) > 1 else ""
            years = list(re.finditer(r"\b((?:19|20)\d{2})\b", after))
            year_match = years[-1] if years else None
            year = year_match.group(1) if year_match else ""
            venue = after[: year_match.start()].strip(" ,.;") if year_match else ""
            tail = after[year_match.end() :].strip(" ,.;") if year_match else ""
            volume_pages = re.search(r"(\d+)(?:\(([^)]+)\))?\s*:\s*(\d+(?:\s*[-–]\s*\d+)?)", tail)
            return {
                "id": f"ref{index}", "type": self._type(marker.group(1)), "authors": authors, "title": title,
                "year": year, "venue": venue, "volume": volume_pages.group(1) if volume_pages else "",
                "issue": volume_pages.group(2) if volume_pages and volume_pages.group(2) else "",
                "pages": volume_pages.group(3) if volume_pages else "", "doi": "", "url": "", "raw": raw,
            }

        year_match = re.search(r"\b((?:19|20)\d{2})\b", text)
        year = year_match.group(1) if year_match else ""
        before = text[: year_match.start()].strip(" .;,(" ) if year_match else ""
        after = text[year_match.end() :].lstrip("). ").strip(" .;,") if year_match else ""
        author_text = before.rsplit(".", 1)[0] if "." in before else before
        authors = [item.strip() for item in re.split(r"\s+(?:and|&|\u548c)\s+|;", author_text) if item.strip()]
        title, _, remainder = after.partition(".")
        title = title.strip()
        venue = re.split(r",\s*\d+(?:\s*\(|\s*,|\s*:)", remainder.strip())[0].strip(" .;,")
        volume_pages = re.search(r"\b(\d+)\s*(?:\(([^)]+)\))?\s*[,;:]\s*(\d+(?:\s*[-–]\s*\d+)?)", remainder)
        doi_match = re.search(r"(?:https?://doi\.org/|doi:\s*)(10\.\S+?)(?:[.\s]|$)", text, re.IGNORECASE)
        return {
            "id": f"ref{index}", "type": "article-journal", "authors": authors, "title": title,
            "year": year, "venue": venue, "volume": volume_pages.group(1) if volume_pages else "",
            "issue": volume_pages.group(2) if volume_pages and volume_pages.group(2) else "",
            "pages": volume_pages.group(3) if volume_pages else "", "doi": doi_match.group(1).rstrip(".") if doi_match else "",
            "url": "", "raw": raw,
        }

    def _quality_gate(self, source: Path, refs: list[dict[str, Any]], reference_index: int | None) -> Path:
        invalid = [
            {"index": index, "missing": [field for field in ("authors", "title", "year", "venue") if not item.get(field)], "raw": item.get("raw", "")}
            for index, item in enumerate(refs, 1)
            if any(not item.get(field) for field in ("authors", "title", "year", "venue"))
        ]
        report = {"source": str(source), "reference_heading_index": reference_index, "reference_count": len(refs), "invalid": invalid, "passed": bool(refs) and not invalid}
        path = self.session_dir / "reference_quality_report.json"
        path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        if not report["passed"]:
            raise ValueError(f"Reference quality gate failed: {len(invalid)} incomplete entries. Report: {path}")
        return path

    def _write_nature_document(self, source: Path, refs: list[dict[str, Any]], reference_index: int | None, nature_markdown: Path, arguments: dict[str, Any], style: str = "nature") -> tuple[Path, Path, str]:
        if reference_index is None or reference_index < 0:
            raise ValueError("Cannot replace a DOCX reference list without a recognized reference heading")
        from docx import Document
        before = self._sha256(source)
        requested = str(arguments.get("output_path") or "").strip()
        output = Path(requested).expanduser() if requested else self.session_dir / "artifacts" / f"{source.stem}_{style}.docx"
        if not output.is_absolute():
            output = self.session_dir / output.name
        if output.resolve() == source.resolve():
            raise ValueError("Output DOCX must not overwrite the source document")
        output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, output)
        document = Document(output)
        if reference_index >= len(document.paragraphs):
            raise ValueError("Reference heading changed while preparing the copied document")
        document.paragraphs[reference_index].text = "References"
        for paragraph in document.paragraphs[reference_index + 1 :]:
            paragraph._element.getparent().remove(paragraph._element)
        audit = self._convert_inline_citations(document, reference_index, len(refs)) if arguments.get("convert_in_text", True) else {"converted": 0, "unresolved": []}
        for line in self._nature_lines(nature_markdown):
            document.add_paragraph(line)
        document.save(output)
        if self._sha256(source) != before:
            raise RuntimeError("Source document changed during conversion; output is not trusted")
        audit_path = self.session_dir / "nature_citation_audit.json"
        audit_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
        return output, audit_path, before

    def _convert_inline_citations(self, document: Any, reference_index: int, reference_count: int) -> dict[str, Any]:
        pattern = re.compile(r"\[(\d+(?:\s*[,\-–]\s*\d+)*)\]")
        converted, unresolved = 0, []
        for paragraph in document.paragraphs[:reference_index]:
            for run in list(paragraph.runs):
                matches = list(pattern.finditer(run.text))
                if not matches:
                    continue
                if any(not self._citation_numbers(match.group(1), reference_count) for match in matches):
                    unresolved.extend(match.group(0) for match in matches)
                    continue
                converted += self._replace_run_citations(run, matches)
        return {"converted": converted, "unresolved": sorted(set(unresolved))}

    def _citation_numbers(self, value: str, reference_count: int) -> bool:
        numbers = [int(item) for item in re.findall(r"\d+", value)]
        return bool(numbers) and all(1 <= item <= reference_count for item in numbers)

    def _replace_run_citations(self, run: Any, matches: list[re.Match[str]]) -> int:
        from docx.oxml import OxmlElement
        original, parent, cursor, inserted = run._r, run._r.getparent(), 0, 0
        for match in matches:
            for text, superscript in ((run.text[cursor : match.start()], False), (match.group(1).replace("-", "–"), True)):
                if not text:
                    continue
                clone = deepcopy(original)
                clone.clear_content()
                text_node = OxmlElement("w:t")
                text_node.text = text
                clone.append(text_node)
                if superscript:
                    vertical = OxmlElement("w:vertAlign")
                    vertical.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "superscript")
                    clone.get_or_add_rPr().append(vertical)
                    inserted += 1
                original.addprevious(clone)
            cursor = match.end()
        if cursor < len(run.text):
            clone = deepcopy(original)
            clone.clear_content()
            text_node = OxmlElement("w:t")
            text_node.text = run.text[cursor:]
            clone.append(text_node)
            original.addprevious(clone)
        parent.remove(original)
        return inserted

    def _nature_lines(self, path: Path) -> list[str]:
        return [line for line in path.read_text(encoding="utf-8", errors="ignore").splitlines() if line and not line.startswith("#")]

    def _sha256(self, path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for block in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(block)
        return digest.hexdigest()

    def _type(self, marker: str) -> str:
        return {"J": "article-journal", "C": "paper-conference", "M": "book", "D": "thesis", "R": "report"}.get(marker.upper(), "article-journal")

    def _source_order_authors(self, values: list[str]) -> list[str]:
        """DOCX [J]/[C] records commonly use 'Surname Initials'; formatter expects 'Initials Surname'."""
        normalized = []
        for raw in values:
            value = raw.strip()
            if not value or value.lower() == "et al":
                if value:
                    normalized.append(value)
                continue
            parts = value.replace(".", "").split()
            if len(parts) >= 2 and all(re.fullmatch(r"[A-Za-z]{1,2}", item) for item in parts[1:]):
                normalized.append(" ".join([*parts[1:], parts[0]]))
            else:
                normalized.append(value)
        return normalized

    def _bibtex(self, text: str) -> list[dict[str, Any]]:
        refs = []
        for match in re.finditer(r"@(\w+)\s*\{\s*([^,]+),(.*?)(?=\n?\s*@|\Z)", text, re.DOTALL):
            entry_type, key, body = match.groups()
            fields = {item.group(1).lower(): re.sub(r"\s+", " ", item.group(2) or item.group(3) or "").strip() for item in re.finditer(r"(\w[\w-]*)\s*=\s*(?:\{(.*?)\}|\"(.*?)\")\s*,?", body, re.DOTALL)}
            refs.append({"id": key.strip(), "type": entry_type.lower(), "title": fields.get("title", ""), "authors": [item.strip() for item in fields.get("author", "").split(" and ") if item.strip()], "year": fields.get("year", ""), "venue": fields.get("journal") or fields.get("booktitle") or "", "volume": fields.get("volume", ""), "issue": fields.get("number", ""), "pages": fields.get("pages", ""), "publisher": fields.get("publisher", ""), "doi": fields.get("doi", ""), "url": fields.get("url", "")})
        if not refs:
            raise ValueError("No BibTeX entries could be parsed")
        return refs

    def _tagged(self, text: str) -> list[dict[str, Any]]:
        records, current = [], {}
        for raw in text.splitlines():
            match = re.match(r"^([A-Z0-9]{2,4})\s*-\s*(.*)$", raw.strip())
            if not match:
                continue
            tag, value = match.groups()
            if tag in {"ER", "EF"}:
                if current:
                    records.append(current)
                    current = {}
            else:
                current.setdefault(tag, []).append(value.strip())
        if current:
            records.append(current)
        refs = []
        for item in records:
            get = lambda *tags: next((item[tag][0] for tag in tags if item.get(tag)), "")
            refs.append({"type": get("TY", "PT").lower(), "title": get("TI", "T1"), "authors": item.get("AU") or item.get("A1") or item.get("FAU") or [], "year": get("PY", "Y1", "DP")[:4], "venue": get("JO", "JF", "T2", "JT"), "volume": get("VL", "VI"), "issue": get("IS", "IP"), "pages": "-".join(filter(None, (get("SP"), get("EP")))), "publisher": get("PB"), "doi": get("DO", "LID").replace(" [doi]", ""), "url": get("UR", "L2")})
        if not refs:
            raise ValueError("No RIS/NBIB entries could be parsed")
        return refs

    def _style(self, value: str) -> str:
        normalized = value.strip().lower().replace("/", "").replace(" ", "-")
        return {"gbt7714": "gbt7714-numeric", "gbt-7714": "gbt7714-numeric", "gbt7714-2015": "gbt7714-numeric", "iop-publishing": "iop"}.get(normalized, normalized)
