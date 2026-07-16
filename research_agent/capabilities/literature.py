from __future__ import annotations

import json
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import replace
from pathlib import Path
from threading import Event
from typing import Any

from research_agent.config import AgentConfig
from research_agent.context import ContextManager
from research_agent.logging import ModelCallLogger
from research_agent.tools.llm_client import LLMClient
from research_agent.tools.arxiv import ArxivClient
from research_agent.tools.crossref import CrossrefClient
from research_agent.tools.openalex import OpenAlexClient
from research_agent.tools.pubmed import PubMedClient
from research_agent.tools.semantic_scholar import SemanticScholarClient


def normalized_title(title: str) -> str:
    return re.sub(r"\W+", " ", (title or "").lower()).strip()


class LiteratureService:
    def __init__(self, config: AgentConfig, session_dir: Path, cancel_event: Event | None = None):
        self.config = config
        self.session_dir = session_dir
        self.raw_dir = session_dir / "raw_search"
        self.raw_dir.mkdir(parents=True, exist_ok=True)
        structured_config = replace(
            config,
            llm_timeout_seconds=config.structured_llm_timeout_seconds,
            llm_max_tokens=config.structured_llm_max_tokens,
            llm_retry=config.structured_llm_retry,
        )
        self.client = LLMClient(structured_config, ModelCallLogger(session_dir), cancel_event)
        self._assessment_cache: dict[str, dict[str, Any]] = {}
        self._last_next_query = ""
        self._last_semantic_stop = False
        self._last_semantic_stop_reason = ""
        self._last_assessment_available = True

    def search(self, arguments: dict[str, Any]) -> dict[str, Any]:
        queries = list(dict.fromkeys(
            str(item).strip()
            for item in ([arguments.get("query")] + list(arguments.get("queries") or []))
            if str(item or "").strip()
        ))
        if not queries:
            raise ValueError("Literature search needs a topic or query")
        queries = queries[:self.config.literature_max_batch_queries]
        query = queries[0]
        if len(queries) > 1:
            return self._search_batch(arguments, queries)
        year_from = self._int(arguments.get("year_from"))
        year_to = self._int(arguments.get("year_to"))
        venues = [str(item).strip() for item in arguments.get("venues") or [] if str(item).strip()]
        sources = self._selected_sources(arguments)
        limit = max(1, min(self.config.literature_max_limit, self._int(arguments.get("limit")) or self.config.literature_default_limit))
        max_rounds = max(1, min(self.config.literature_max_rounds, self._int(arguments.get("rounds")) or 1))
        max_requests_per_source = max(1, min(
            self.config.literature_max_requests_per_source,
            self._int(arguments.get("max_requests_per_source")) or self.config.literature_max_requests_per_source,
        ))
        must_include = [str(item) for item in arguments.get("must_include") or [] if str(item).strip()]
        exclude = [str(item) for item in arguments.get("exclude") or [] if str(item).strip()]
        precise = bool(arguments.get("precise") or must_include or exclude)
        if arguments.get("no_network"):
            raise ValueError("当前请求禁止联网；请先上传资料，或取消不要联网限制")

        pool: list[dict[str, Any]] = []
        errors: list[dict[str, str]] = []
        rounds: list[dict[str, Any]] = []
        current_query = query
        source_attempts: dict[str, int] = {source: 0 for source in sources}
        blocked_sources: set[str] = set()

        for round_number in range(1, max_rounds + 1):
            before = len(pool)
            round_results: list[dict[str, Any]] = []
            source_requests: list[dict[str, Any]] = []
            for source in sources:
                if source in blocked_sources:
                    source_requests.append({"source": source, "status": "skipped_blocked"})
                    continue
                if source_attempts.get(source, 0) >= max_requests_per_source:
                    source_requests.append({"source": source, "status": "skipped_budget"})
                    continue
                client = self._client(source)
                if client is None:
                    errors.append({"source": source, "error": "unsupported source"})
                    source_requests.append({"source": source, "status": "unsupported"})
                    continue
                try:
                    source_attempts[source] = source_attempts.get(source, 0) + 1
                    fetch_limit = min(max(limit, 10), self.config.literature_screening_limit)
                    papers = client.search(
                        current_query,
                        year_from=year_from,
                        year_to=year_to,
                        limit=fetch_limit,
                    )
                    (self.raw_dir / f"round_{round_number}_{source}.json").write_text(
                        json.dumps(papers, ensure_ascii=False, indent=2), encoding="utf-8"
                    )
                    round_results.extend(papers)
                    source_requests.append({"source": source, "status": "ok", "count": len(papers), "attempt": source_attempts[source]})
                    if self.config.literature_request_delay_seconds:
                        time.sleep(self.config.literature_request_delay_seconds)
                except Exception as exc:
                    error = str(exc)
                    errors.append({"source": source, "error": error})
                    source_requests.append({"source": source, "status": "error", "error": error, "attempt": source_attempts.get(source, 0)})
                    if "429" in error or "too many" in error.lower() or "rate" in error.lower():
                        blocked_sources.add(source)
            pool = self._dedupe([*pool, *round_results])
            candidates = pool[:self.config.literature_screening_limit]
            screened, excluded, edge = self.screen_with_edge(candidates, query, year_from, year_to, venues, must_include, exclude, precise)
            next_query = self._last_next_query.strip()
            new_terms = [next_query] if next_query and next_query != current_query else []
            stop_reason = ""
            if len(screened) >= limit:
                stop_reason = "target_reached"
            elif len(pool) == before:
                stop_reason = "no_new_papers"
            elif not any(item.get("status") == "ok" for item in source_requests):
                stop_reason = "all_sources_failed_or_blocked"
            elif not new_terms:
                stop_reason = "no_new_terms"
            elif self._last_semantic_stop:
                stop_reason = self._last_semantic_stop_reason or "semantic_saturation"
            rounds.append(
                {
                    "round": round_number,
                    "query": current_query,
                    "new_papers": len(pool) - before,
                    "pool_size": len(pool),
                    "screened": len(screened),
                    "edge": len(edge),
                    "new_terms": new_terms,
                    "source_requests": source_requests,
                    "stop_reason": stop_reason or ("max_rounds" if round_number == max_rounds else "continue"),
                }
            )
            if stop_reason or round_number == max_rounds:
                break
            current_query = next_query

        candidates = pool[:self.config.literature_screening_limit]
        screened, excluded, edge = self.screen_with_edge(candidates, query, year_from, year_to, venues, must_include, exclude, precise)
        screened = screened[:limit]
        outcome = self._search_outcome(screened, errors, [request for item in rounds for request in item["source_requests"]])
        paths = self._write_outputs(query, screened, excluded, edge, rounds, errors)
        requested = str(arguments.get("output_path") or "").strip()
        if requested:
            paths["requested_output"] = str(self._save_requested(requested, query, screened))
        return {
            "message": self._search_message(screened, errors, paths, outcome),
            "artifacts": paths,
            "data": {"count": len(screened), "papers": screened, "rounds": rounds, "errors": errors},
            "model_data": {**self._answer_evidence(query, screened, errors), "outcome": outcome},
            "outcome": outcome,
            "progress": {
                "summary": self._progress_summary(outcome, retrieved=len(pool), included=len(screened), errors=errors),
                "metrics": {"retrieved": len(pool), "evaluated": len(candidates), "included": len(screened), "excluded": len(excluded), "edge": len(edge)},
            },
        }

    def _search_batch(self, arguments: dict[str, Any], queries: list[str]) -> dict[str, Any]:
        year_from = self._int(arguments.get("year_from"))
        year_to = self._int(arguments.get("year_to"))
        venues = [str(item).strip() for item in arguments.get("venues") or [] if str(item).strip()]
        sources = self._selected_sources(arguments)
        limit = max(1, min(self.config.literature_max_limit, self._int(arguments.get("limit")) or self.config.literature_default_limit))
        must_include = [str(item) for item in arguments.get("must_include") or [] if str(item).strip()]
        exclude = [str(item) for item in arguments.get("exclude") or [] if str(item).strip()]
        precise = bool(arguments.get("precise") or must_include or exclude)
        if arguments.get("no_network"):
            raise ValueError("当前请求禁止联网；请先上传资料，或取消不要联网限制。")

        results: list[dict[str, Any]] = []
        result_batches: dict[tuple[int, str], list[dict[str, Any]]] = {}
        errors: list[dict[str, str]] = []
        requests: list[dict[str, Any]] = []

        def fetch(query_index: int, source: str, scholarly_query: str) -> tuple[int, str, str, list[dict[str, Any]]]:
            client = self._client(source)
            if client is None:
                raise ValueError("unsupported source")
            fetch_limit = min(max(limit, 10), self.config.literature_screening_limit)
            papers = client.search(scholarly_query, year_from=year_from, year_to=year_to, limit=fetch_limit)
            return query_index, source, scholarly_query, papers

        jobs = [(index, source, scholarly_query) for index, scholarly_query in enumerate(queries, 1) for source in sources]
        with ThreadPoolExecutor() as pool:
            futures = {pool.submit(fetch, *job): job for job in jobs}
            for future in as_completed(futures):
                query_index, source, scholarly_query = futures[future]
                try:
                    _, _, _, papers = future.result()
                    (self.raw_dir / f"batch_{query_index}_{source}.json").write_text(
                        json.dumps(papers, ensure_ascii=False, indent=2), encoding="utf-8"
                    )
                    result_batches[(query_index, source)] = papers
                    requests.append({"query": scholarly_query, "source": source, "status": "ok", "count": len(papers)})
                except Exception as exc:
                    error = str(exc)
                    errors.append({"source": source, "query": scholarly_query, "error": error})
                    requests.append({"query": scholarly_query, "source": source, "status": "error", "error": error})

        ordered_batches = [
            result_batches.get((index, source), [])
            for index, _query in enumerate(queries, 1)
            for source in sources
        ]
        for offset in range(max((len(batch) for batch in ordered_batches), default=0)):
            for batch in ordered_batches:
                if offset < len(batch):
                    results.append(batch[offset])
        deduplicated = self._dedupe(results)
        screening_candidates = deduplicated[:min(limit, self.config.literature_screening_limit)]
        research_question = str(arguments.get("request") or queries[0]).strip()
        screened, excluded, edge = self.screen_with_edge(
            screening_candidates, research_question, year_from, year_to, venues, must_include, exclude, precise
        )
        screened = screened[:limit]
        rounds = [{
            "round": 1,
            "queries": queries,
            "pool_size": len(deduplicated),
            "screened": len(screened),
            "edge": len(edge),
            "source_requests": requests,
            "stop_reason": "batch_complete",
        }]
        outcome = self._search_outcome(screened, errors, requests)
        paths = self._write_outputs(research_question, screened, excluded, edge, rounds, errors)
        requested = str(arguments.get("output_path") or "").strip()
        if requested:
            paths["requested_output"] = str(self._save_requested(requested, research_question, screened))
        return {
            "message": self._search_message(screened, errors, paths, outcome),
            "artifacts": paths,
            "data": {"count": len(screened), "papers": screened, "rounds": rounds, "errors": errors},
            "model_data": {**self._answer_evidence(research_question, screened, errors), "outcome": outcome},
            "outcome": outcome,
            "progress": {
                "summary": self._progress_summary(outcome, retrieved=len(results), included=len(screened), errors=errors),
                "metrics": {"retrieved": len(results), "deduplicated": len(deduplicated), "evaluated": len(screening_candidates), "included": len(screened), "excluded": len(excluded), "edge": len(edge)},
            },
        }

    def screen_active(self, papers_path: str, arguments: dict[str, Any]) -> dict[str, Any]:
        papers = json.loads(Path(papers_path).read_text(encoding="utf-8-sig"))
        query = str(arguments.get("query") or arguments.get("topic") or "").strip()
        if not query:
            query = str(arguments.get("keywords") or "").strip()
        if not query:
            raise ValueError("Screening needs a query")
        screened, excluded, edge = self.screen_with_edge(
            papers,
            query,
            self._int(arguments.get("year_from")),
            self._int(arguments.get("year_to")),
            [str(item) for item in arguments.get("venues") or []],
            [str(item) for item in arguments.get("must_include") or []],
            [str(item) for item in arguments.get("exclude") or []],
            bool(arguments.get("precise") or arguments.get("must_include") or arguments.get("exclude")),
        )
        active = self.session_dir / "screened_papers.json"
        dropped = self.session_dir / "excluded_papers.json"
        edge_path = self.session_dir / "edge_papers.json"
        active.write_text(json.dumps(screened, ensure_ascii=False, indent=2), encoding="utf-8")
        dropped.write_text(json.dumps(excluded, ensure_ascii=False, indent=2), encoding="utf-8")
        edge_path.write_text(json.dumps(edge, ensure_ascii=False, indent=2), encoding="utf-8")
        return {
            "message": f"筛选完成：保留 {len(screened)} 篇，边缘 {len(edge)} 篇，排除 {len(excluded)} 篇。",
            "artifacts": {"active_papers": str(active), "excluded_papers": str(dropped), "edge_papers": str(edge_path)},
            "data": {"count": len(screened), "papers": screened, "edge": edge},
            "model_data": self._answer_evidence(query, screened, []),
        }

    def screen(
        self,
        papers: list[dict[str, Any]],
        query: str,
        year_from: int | None,
        year_to: int | None,
        venues: list[str],
    ) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        included, excluded, edge = self.screen_with_edge(papers, query, year_from, year_to, venues, [], [], False)
        return included, [*edge, *excluded]

    def screen_with_edge(
        self,
        papers: list[dict[str, Any]],
        query: str,
        year_from: int | None,
        year_to: int | None,
        venues: list[str],
        must_include: list[str],
        exclude: list[str],
        precise: bool,
    ) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
        included: list[dict[str, Any]] = []
        excluded: list[dict[str, Any]] = []
        edge: list[dict[str, Any]] = []
        eligible: list[tuple[str, dict[str, Any]]] = []
        for index, raw in enumerate(papers):
            paper = dict(raw)
            year = self._int(paper.get("year"))
            venue = str(paper.get("venue") or "")
            reason = ""
            if year_from and year and year < year_from:
                reason = "before_year_range"
            elif year_to and year and year > year_to:
                reason = "after_year_range"
            elif venues and not any(item.lower() in venue.lower() for item in venues):
                reason = "venue_not_matched"
            if reason:
                paper["exclusion_reason"] = reason
                paper["screening_reasons"] = [reason]
                excluded.append(paper)
                continue
            eligible.append((f"p{index}", paper))

        assessment = self._semantic_assessment(
            eligible, query=query, must_include=must_include, exclude=exclude, precise=precise
        )
        decisions = {
            str(item.get("id") or ""): item
            for item in assessment.get("papers", [])
            if isinstance(item, dict)
        }
        self._last_next_query = str(assessment.get("next_query") or "").strip()
        self._last_semantic_stop = bool(assessment.get("stop"))
        self._last_semantic_stop_reason = str(assessment.get("stop_reason") or "")
        self._last_assessment_available = bool(assessment.get("_assessment_available", True))
        for paper_id, paper in eligible:
            decision = decisions.get(paper_id, {})
            bucket = str(decision.get("decision") or "edge").lower()
            if bucket not in {"include", "edge", "exclude"}:
                bucket = "edge"
            try:
                score = max(0.0, min(100.0, float(decision.get("score"))))
            except (TypeError, ValueError):
                score = None
            reasons = decision.get("reasons") if isinstance(decision.get("reasons"), list) else []
            concise_reason = str(decision.get("reason") or "").strip()
            if not reasons and concise_reason:
                reasons = [concise_reason]
            paper["relevance_score"] = score
            paper["screening_reasons"] = [str(item) for item in reasons if str(item).strip()]
            paper["matched_required_terms"] = [
                str(item) for item in (decision.get("matched_requirements") or []) if str(item).strip()
            ]
            paper["missing_required_terms"] = [
                str(item) for item in (decision.get("missing_requirements") or []) if str(item).strip()
            ]
            paper["semantic_assessment"] = "model" if decision else "unavailable"
            if bucket == "exclude":
                paper["exclusion_reason"] = str(decision.get("exclusion_reason") or "semantic_exclusion")
                excluded.append(paper)
            elif bucket == "edge":
                paper["exclusion_reason"] = str(decision.get("exclusion_reason") or "needs_human_review")
                edge.append(paper)
            else:
                included.append(paper)

        def rank(item: dict[str, Any]) -> tuple[float, int]:
            score = item.get("relevance_score")
            return (
                float(score) if isinstance(score, (int, float)) else -1.0,
                self._int(item.get("citation_count")) or 0,
            )

        included.sort(key=rank, reverse=True)
        excluded.sort(key=rank, reverse=True)
        edge.sort(key=rank, reverse=True)
        return included, excluded, edge

    def _semantic_assessment(
        self,
        papers: list[tuple[str, dict[str, Any]]],
        *,
        query: str,
        must_include: list[str],
        exclude: list[str],
        precise: bool,
    ) -> dict[str, Any]:
        candidates = [
            {
                "id": paper_id,
                "title": paper.get("title"),
                "abstract": str(paper.get("abstract") or ""),
                "year": paper.get("year"),
                "venue": paper.get("venue"),
                "keywords": paper.get("keywords"),
                "source": paper.get("source"),
                "type": paper.get("type") or (paper.get("raw_metadata") or {}).get("type"),
            }
            for paper_id, paper in papers
        ]
        candidates = self._fit_screening_candidates(candidates)
        cache_key = json.dumps(
            {"query": query, "must_include": must_include, "exclude": exclude, "precise": precise, "papers": candidates},
            ensure_ascii=False,
            sort_keys=True,
            default=str,
        )
        if cache_key in self._assessment_cache:
            return self._assessment_cache[cache_key]
        if not candidates:
            return {
                "papers": [], "next_query": "", "stop": True,
                "stop_reason": "no_eligible_candidates", "_assessment_available": False,
            }
        prompt = json.dumps(
            {
                "research_question": query,
                "required_concepts": must_include,
                "excluded_concepts_or_study_types": exclude,
                "strict_all_requirements": precise,
                "candidates": candidates,
                "output_contract": {
                    "included": [{"id": "candidate id", "score": "0..100", "reason": "one short sentence"}],
                    "edge": [{"id": "candidate id", "score": "0..100", "reason": "one short sentence"}],
                    "excluded_ids": ["candidate id"],
                    "next_query": "a complete refined scholarly query, or empty string",
                    "stop": "boolean", "stop_reason": "short reason",
                },
            },
            ensure_ascii=False,
        )
        result = self.client.complete(
            "literature_semantic_screening",
            prompt,
            system=(
                "You are a conservative academic screening evaluator. Interpret scientific meaning, negation, synonyms, "
                "population, method, and study type from supplied metadata only. Return one valid JSON object matching "
                "the contract and no prose. Give reasons only for included or edge candidates; excluded candidates go "
                "only in excluded_ids. Never invent missing evidence. Borderline or insufficient evidence must be edge."
            ),
            temperature=0,
        )
        parsed = self._json_object(result.text) if result.used_remote_model else None
        if isinstance(parsed, dict) and not isinstance(parsed.get("papers"), list):
            compact_shape = all(isinstance(parsed.get(key, []), list) for key in ("included", "edge", "excluded_ids"))
            if compact_shape and any(key in parsed for key in ("included", "edge", "excluded_ids")):
                normalized: list[dict[str, Any]] = []
                for decision, key in (("include", "included"), ("edge", "edge")):
                    for item in parsed.get(key, []):
                        entry = dict(item) if isinstance(item, dict) else {"id": str(item)}
                        entry["decision"] = decision
                        normalized.append(entry)
                normalized.extend(
                    {"id": str(item), "decision": "exclude"}
                    for item in parsed.get("excluded_ids", [])
                )
                parsed["papers"] = normalized
        if isinstance(parsed, dict) and isinstance(parsed.get("papers"), list):
            parsed["_assessment_available"] = True
        else:
            parsed = {
                "papers": [
                    {
                        "id": paper_id, "decision": "edge", "score": None,
                        "reason": "semantic_assessment_unavailable",
                        "exclusion_reason": "needs_human_review",
                    }
                    for paper_id, _paper in papers
                ],
                "next_query": "", "stop": True, "stop_reason": "semantic_assessment_unavailable",
                "_assessment_available": False,
            }
        self._assessment_cache[cache_key] = parsed
        return parsed

    def _fit_screening_candidates(self, candidates: list[dict[str, Any]], budget: int | None = None) -> list[dict[str, Any]]:
        """Fit candidate evidence to the configured observation budget without topic rules."""
        budget = max(256, int(budget or self.config.context_observation_budget))
        if ContextManager.estimate_tokens(candidates) <= budget:
            return candidates
        longest = max((len(str(item.get("abstract") or "")) for item in candidates), default=0)
        low, high = 0, longest
        best = [{**item, "abstract": ""} for item in candidates]
        while low <= high:
            width = (low + high) // 2
            probe = [{**item, "abstract": str(item.get("abstract") or "")[:width]} for item in candidates]
            if ContextManager.estimate_tokens(probe) <= budget:
                best = probe
                low = width + 1
            else:
                high = width - 1
        return best

    def _answer_evidence(
        self,
        research_question: str,
        papers: list[dict[str, Any]],
        errors: list[dict[str, Any]],
    ) -> dict[str, Any]:
        """Provide a compact, sufficient synthesis payload instead of forcing artifact rereads."""
        evidence = [
            {
                "id": f"p{index}",
                "title": paper.get("title"),
                "authors": list(paper.get("authors") or [])[:3],
                "year": paper.get("year"),
                "venue": paper.get("venue"),
                "abstract": str(paper.get("abstract") or ""),
                "doi": paper.get("doi"),
                "url": paper.get("url") or paper.get("open_access_url"),
                "relevance_score": paper.get("relevance_score"),
                "screening_reasons": paper.get("screening_reasons") or [],
            }
            for index, paper in enumerate(papers, 1)
        ]
        evidence_budget = max(256, int(self.config.context_observation_budget * 0.72))
        answer_ready = bool(papers) and self._last_assessment_available
        return {
            "answer_ready": answer_ready,
            "missing_evidence": [] if answer_ready else ["No semantically screened papers are available."],
            "research_question": research_question,
            "included_count": len(papers),
            "papers": self._fit_screening_candidates(evidence, evidence_budget),
            "source_errors": errors,
            "completion_guidance": (
                "The supplied evidence is sufficient for a concise user-facing synthesis. "
                "Answer now unless the user explicitly requested an additional artifact or a named missing field."
            ) if answer_ready else "Explain the evidence gap plainly; do not present unscreened candidates as findings.",
        }

    @staticmethod
    def _json_object(text: str) -> dict[str, Any] | None:
        value = str(text or "").strip()
        if value.startswith("```"):
            lines = value.splitlines()
            value = "\n".join(lines[1:-1]).strip() if len(lines) >= 3 else ""
        try:
            parsed = json.loads(value)
        except json.JSONDecodeError:
            return None
        return parsed if isinstance(parsed, dict) else None

    def _dedupe(self, papers: list[dict[str, Any]]) -> list[dict[str, Any]]:
        merged: dict[str, dict[str, Any]] = {}
        for paper in papers:
            key = str(paper.get("doi") or "").lower().strip() or normalized_title(str(paper.get("title") or ""))
            if not key:
                continue
            if key not in merged:
                item = dict(paper)
                item["sources"] = [paper.get("source")] if paper.get("source") else []
                merged[key] = item
                continue
            current = merged[key]
            if len(str(paper.get("abstract") or "")) > len(str(current.get("abstract") or "")):
                current["abstract"] = paper.get("abstract")
            current["citation_count"] = max(
                self._int(current.get("citation_count")) or 0,
                self._int(paper.get("citation_count")) or 0,
            )
            for field in ("doi", "url", "venue", "year", "open_access_url"):
                if not current.get(field) and paper.get(field):
                    current[field] = paper[field]
            source = paper.get("source")
            if source and source not in current["sources"]:
                current["sources"].append(source)
        return list(merged.values())

    def _write_outputs(
        self,
        query: str,
        screened: list[dict[str, Any]],
        excluded: list[dict[str, Any]],
        edge: list[dict[str, Any]],
        rounds: list[dict[str, Any]],
        errors: list[dict[str, str]],
    ) -> dict[str, str]:
        paths = {
            "raw_papers": self.session_dir / "papers_raw.json",
            "active_papers": self.session_dir / "screened_papers.json",
            "excluded_papers": self.session_dir / "excluded_papers.json",
            "edge_papers": self.session_dir / "edge_papers.json",
            "recursive_search_plan": self.session_dir / "recursive_search_plan.json",
            "paper_pool_markdown": self.session_dir / "paper_pool.md",
        }
        paths["raw_papers"].write_text(json.dumps(self._dedupe([*screened, *edge, *excluded]), ensure_ascii=False, indent=2), encoding="utf-8")
        paths["active_papers"].write_text(json.dumps(screened, ensure_ascii=False, indent=2), encoding="utf-8")
        paths["excluded_papers"].write_text(json.dumps(excluded, ensure_ascii=False, indent=2), encoding="utf-8")
        paths["edge_papers"].write_text(json.dumps(edge, ensure_ascii=False, indent=2), encoding="utf-8")
        paths["recursive_search_plan"].write_text(
            json.dumps({"query": query, "rounds": rounds, "errors": errors}, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        paths["paper_pool_markdown"].write_text(self._markdown(query, screened, edge), encoding="utf-8")
        if errors:
            error_path = self.session_dir / "literature_search_errors.json"
            error_path.write_text(json.dumps(errors, ensure_ascii=False, indent=2), encoding="utf-8")
            paths["literature_search_errors"] = error_path
        return {key: str(path) for key, path in paths.items()}

    def _markdown(self, query: str, papers: list[dict[str, Any]], edge: list[dict[str, Any]] | None = None) -> str:
        edge = edge or []
        lines = ["# Screened Literature Pool", "", f"Query: {query}", "", "## Included", ""]
        for index, paper in enumerate(papers, 1):
            title = paper.get("title") or "Untitled"
            url = paper.get("url") or paper.get("open_access_url") or ""
            title_text = f"[{title}]({url})" if url else title
            lines.extend(
                [
                    f"{index}. {title_text}",
                    f"   - Year: {paper.get('year') or 'unknown'}",
                    f"   - Venue: {paper.get('venue') or 'unknown'}",
                    f"   - Score: {paper.get('relevance_score', 0)}",
                    f"   - DOI: {paper.get('doi') or 'not available'}",
                ]
            )
        if edge:
            lines.extend(["", "## Borderline / Edge Papers", ""])
            for index, paper in enumerate(edge, 1):
                lines.extend(
                    [
                        f"{index}. {paper.get('title') or 'Untitled'}",
                        f"   - Year: {paper.get('year') or 'unknown'}",
                        f"   - Venue: {paper.get('venue') or 'unknown'}",
                        f"   - Reason: {paper.get('exclusion_reason') or 'borderline'}",
                    ]
                )
        return "\n".join(lines) + "\n"

    def _save_requested(self, raw: str, query: str, papers: list[dict[str, Any]]) -> Path:
        path = Path(raw).expanduser()
        if not path.suffix:
            path = path / "literature_search.md"
        path.parent.mkdir(parents=True, exist_ok=True)
        if path.suffix.lower() == ".docx":
            try:
                from docx import Document
            except ImportError as exc:
                raise RuntimeError("DOCX output requires python-docx") from exc
            document = Document()
            document.add_heading("Screened Literature Pool", level=1)
            document.add_paragraph(f"Query: {query}")
            for index, paper in enumerate(papers, 1):
                document.add_heading(f"{index}. {paper.get('title') or 'Untitled'}", level=2)
                document.add_paragraph(
                    f"Year: {paper.get('year') or 'unknown'}; Venue: {paper.get('venue') or 'unknown'}; "
                    f"DOI: {paper.get('doi') or 'not available'}"
                )
            document.save(path)
        elif path.suffix.lower() == ".json":
            path.write_text(json.dumps(papers, ensure_ascii=False, indent=2), encoding="utf-8")
        else:
            if path.suffix.lower() not in {".md", ".txt"}:
                path = path.with_suffix(".md")
            path.write_text(self._markdown(query, papers), encoding="utf-8")
        return path

    def _search_message(
        self,
        papers: list[dict[str, Any]],
        errors: list[dict[str, str]],
        paths: dict[str, str],
        outcome: str,
    ) -> str:
        if not papers:
            detail = "；".join(f"{item['source']}: {item['error']}" for item in errors[:3])
            lead = {
                "rate_limited": "文献来源当前限流，尚未获得可用论文。",
                "failed": "本次文献检索未成功，尚未获得可用论文。",
                "empty": "已查询所选来源，但没有发现符合条件的论文。",
            }.get(outcome, "没有检索到符合条件的论文。")
            return lead + (f" 来源反馈：{detail}" if detail else " 可以调整关键词、年份或来源后再试。")
        lines = [f"找到并筛选出 {len(papers)} 篇论文："]
        for index, paper in enumerate(papers[:8], 1):
            lines.append(f"{index}. {paper.get('title')} ({paper.get('year') or '?'}, {paper.get('venue') or 'venue unknown'})")
        if len(papers) > 8:
            lines.append(f"其余 {len(papers) - 8} 篇已写入文献池。")
        if errors:
            lines.append("部分数据源失败，但已保留其他来源结果。")
        lines.append(f"文献池：{paths.get('requested_output') or paths['paper_pool_markdown']}")
        return "\n".join(lines)

    @staticmethod
    def _is_rate_limit(error: str) -> bool:
        value = str(error or "").lower()
        return "429" in value or "too many" in value or "rate limit" in value or "rate_limited" in value

    def _search_outcome(
        self,
        papers: list[dict[str, Any]],
        errors: list[dict[str, str]],
        requests: list[dict[str, Any]],
    ) -> str:
        if papers:
            return "partial" if errors else "success"
        successful_requests = any(item.get("status") == "ok" for item in requests)
        if not errors or successful_requests:
            return "empty"
        if errors and all(self._is_rate_limit(str(item.get("error") or "")) for item in errors):
            return "rate_limited"
        return "failed"

    @staticmethod
    def _progress_summary(
        outcome: str,
        *,
        retrieved: int,
        included: int,
        errors: list[dict[str, str]],
    ) -> str:
        if outcome == "success":
            return f"找到 {retrieved} 篇候选文献，筛选后纳入 {included} 篇"
        if outcome == "partial":
            return f"部分来源可用：纳入 {included} 篇，另有 {len(errors)} 个来源请求失败"
        if outcome == "empty":
            return "所选来源已返回，但没有发现符合条件的论文"
        if outcome == "rate_limited":
            return "所选文献来源当前限流，未获得可用结果"
        return "所选文献来源请求失败，未获得可用结果"

    @staticmethod
    def _selected_sources(arguments: dict[str, Any]) -> list[str]:
        sources = list(dict.fromkeys(
            str(item).strip().lower()
            for item in arguments.get("sources") or []
            if str(item).strip()
        ))
        if not sources:
            raise ValueError("多源文献检索需要模型明确选择至少一个 sources 来源")
        return sources

    def _client(self, source: str):
        return {
            "openalex": OpenAlexClient(self.config),
            "pubmed": PubMedClient(self.config),
            "semantic_scholar": SemanticScholarClient(self.config),
            "arxiv": ArxivClient(),
            "crossref": CrossrefClient(self.config),
        }.get(source)

    def _int(self, value: Any) -> int | None:
        try:
            return int(value) if value not in (None, "") else None
        except (TypeError, ValueError):
            return None
